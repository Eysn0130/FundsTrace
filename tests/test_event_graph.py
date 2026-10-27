"""Default business projections count confirmed events; retain ledger evidence."""
from app import event_resolution
from test_event_import import csv_rows, transaction
from test_folder_evidence import accepted


def pair(reference, amount="100.00", *, day="01", currency="CNY"):
    return [transaction(own, **{"交易流水号": reference + "-" + own,
        "跨账簿事件参考号": reference, "交易金额": amount, "交易币种": currency,
        "交易时间": f"2026-01-{day} 10:00:00"}) for own in ("SYN-A", "SYN-B")]


def imported(client, case_id, rows):
    accepted(client, case_id, [("synthetic-events.csv", csv_rows(rows))], scope_bank=None)
    response = client.get(f"/api/v1/cases/{case_id}/flows?granularity=observations")
    assert response.status_code == 200, response.text
    return response.json()["items"]


def track(client, case_id, rows):
    response = client.post("/api/track", json={"case_id": case_id, "flow_ids": [row["id"] for row in rows]})
    assert response.status_code == 200, response.text


def graph(client, case_id):
    response = client.get("/api/graph", params={"case_id": case_id})
    assert response.status_code == 200, response.text
    return response.json()


def projection(client, case_id, **window):
    response = client.post("/api/canvas/investigation/time", json={"case_id": case_id, **window})
    assert response.status_code == 200, response.text
    return response.json()


def test_graph_counts_selected_event_once_and_never_adds_unselected_events(client, make_case):
    cid = make_case("Synthetic selected events")
    rows = imported(client, cid, pair("SYN-E1") + pair("SYN-E2", "99.00", day="02"))
    selected = [row for row in rows if row["trade_no"] == "SYN-E1-SYN-A"]
    track(client, cid, selected)
    edge = graph(client, cid)["edges"][0]
    assert edge["cnt"] == 1 and edge["total"] == "100.00"
    assert edge["event_summary"] == {"confirmed_events": 1, "confirmed_entries": 1,
        "conflict_entries": 0, "unknown_entries": 0,
        "confirmed_amount_totals": [{"currency": "CNY", "amount": "100.00"}]}
    track(client, cid, [row for row in rows if row["trade_no"] == "SYN-E1-SYN-B"])
    edge = graph(client, cid)["edges"][0]
    assert edge["cnt"] == 1 and edge["amount_totals"] == [{"currency": "CNY", "amount": "100.00"}]
    assert edge["event_summary"]["confirmed_events"] == 1 and edge["event_summary"]["confirmed_entries"] == 2
    assert edge["event_summary"]["confirmed_amount_totals"] == [{"currency": "CNY", "amount": "100.00"}]
    assert projection(client, cid)["edges"][0]["event_summary"] == edge["event_summary"]
    payer = next(node for node in graph(client, cid)["nodes"] if node["out_cnt"] == 1)
    assert payer["out_amt"] == "100.00"


def test_time_projection_scopes_event_totals_and_separates_unknown_conflict_and_currencies(client, make_case):
    cid = make_case("Synthetic projected event totals")
    unknown = transaction("SYN-A", **{"交易流水号": "SYN-UNKNOWN", "交易金额": "7.00",
        **{header: "" for header in event_resolution.HEADERS.values()}})
    conflict = pair("SYN-CONFLICT", "8.00")
    conflict[1]["交易金额"] = "9.00"
    rows = imported(client, cid, pair("SYN-E1") + pair("SYN-E2", "99.00", day="02") +
                    pair("SYN-USD", "0.10", currency="USD") + [unknown] + conflict)
    track(client, cid, rows)
    full = graph(client, cid)
    first = projection(client, cid, start="2026-01-01T00:00:00", end="2026-01-01T23:59:59")
    edge = first["edges"][0]
    assert edge["cnt"] is None and edge["total"] is None
    assert edge["funds"]["count"] == {"lower": 2, "upper": 5}
    assert edge["funds"]["totals"] == [
        {"currency": "CNY", "amount": None, "lower": "100.00", "upper": "124.00"},
        {"currency": "USD", "amount": "0.10", "lower": "0.10", "upper": "0.10"}]
    assert edge["event_summary"] == {"confirmed_events": 2, "confirmed_entries": 4,
        "conflict_entries": 2, "unknown_entries": 1,
        "confirmed_amount_totals": [{"currency": "CNY", "amount": "100.00"}, {"currency": "USD", "amount": "0.10"}]}
    second = projection(client, cid, start="2026-01-02T00:00:00", end="2026-01-02T23:59:59")
    assert second["edges"][0]["amount_totals"] == [{"currency": "CNY", "amount": "99.00"}]
    assert second["edges"][0]["event_summary"]["confirmed_amount_totals"] == [{"currency": "CNY", "amount": "99.00"}]
    assert full["edges"][0]["event_summary"]["confirmed_events"] == 3
    assert full["analysis_revision"] == first["analysis_revision"] == second["analysis_revision"]


def test_ordinary_graph_has_no_event_claim_and_case_scope_excludes_foreign_ids(client, make_case):
    cid, other = make_case("Synthetic ordinary graph"), make_case("Synthetic unrelated graph")
    ordinary = [{key: value for key, value in row.items() if key not in event_resolution.HEADERS.values()}
                for row in pair("SYN-ORDINARY")]
    rows = imported(client, cid, ordinary)
    track(client, cid, rows)
    edge = graph(client, cid)["edges"][0]
    assert edge["cnt"] is None and edge["total"] is None and "event_summary" not in edge
    assert edge["funds"]["count"] == {"lower": 0, "upper": 2}
    assert edge["funds"]["totals"] == [{"currency": "CNY", "amount": None, "lower": "0.00", "upper": "200.00"}]
    assert "event_summary" not in projection(client, cid)["edges"][0]
    # A same-case projection now exists, but these selected rows have no
    # declaration. Neither unselected nor foreign confirmed events may leak.
    imported(client, cid, pair("SYN-UNTRACKED", "3.00", day="02"))
    foreign = imported(client, other, pair("SYN-E1", "999.00"))
    edge["ids"] += [row["id"] for row in foreign]
    edge["event_summary"] = {"stale": True}
    event_resolution.decorate_edges(cid, [edge])
    assert "event_summary" not in edge and edge["total"] is None


def test_event_overlay_withdraws_on_conflict_and_restores_after_batch_deletion(client, make_case):
    cid = make_case("Synthetic graph event withdrawal")
    rows = imported(client, cid, pair("SYN-E1"))
    track(client, cid, rows)
    assert graph(client, cid)["edges"][0]["event_summary"]["confirmed_events"] == 1
    _, conflict = accepted(client, cid, [("synthetic-conflict.csv", csv_rows(pair("SYN-E1", "101.00")[:1]))], scope_bank=None)
    for result in (graph(client, cid), projection(client, cid)):
        edge = result["edges"][0]
        assert edge["cnt"] is None and edge["total"] is None
        assert edge["funds"]["totals"] == [{"currency": "CNY", "amount": None, "lower": "0.00", "upper": "200.00"}]
        assert edge["event_summary"] == {"confirmed_events": 0, "confirmed_entries": 0,
            "conflict_entries": 2, "unknown_entries": 0, "confirmed_amount_totals": []}
    response = client.delete(f"/api/batches/{conflict['batch_id']}", params={"case_id": cid})
    assert response.status_code == 200, response.text
    assert projection(client, cid)["edges"][0]["event_summary"]["confirmed_amount_totals"] == [{"currency": "CNY", "amount": "100.00"}]


def test_explicit_member_query_chunks_ids_and_preserves_exact_case_scope(client, database, make_case, monkeypatch):
    cid = make_case("Synthetic bounded member query")
    rows = imported(client, cid, pair("SYN-E1"))
    actual = {row["id"] for row in rows}
    requested = [*range(100000, 104100), *actual, *actual]
    query, calls, edge_calls = database.q_raw, [], []

    def bounded(sql, args=None):
        if "from event_members where" in sql:
            calls.append(len(args) - 1)
            assert len(args) <= 4097
        if "from flows f left join event_members" in sql:
            edge_calls.append(len(args[1]))
            assert len(args[1]) <= 4096
        return query(sql, args)

    monkeypatch.setattr(database, "q_raw", bounded)
    members = event_resolution.members(cid, requested)
    assert {row["flow_id"] for row in members} == actual and len(members) == 2
    assert calls == [4096, 6]
    edge = {"source": rows[0]["payer_account"], "target": rows[0]["payee_account"],
            "ids": requested, "cnt": 2, "total": "200.00"}
    event_resolution.decorate_edges(cid, [edge])
    assert edge_calls == [4096, 6]
    assert edge["event_summary"]["confirmed_entries"] == 2
    assert edge["event_summary"]["unknown_entries"] == 0
    assert edge["cnt"] == 2 and edge["total"] == "200.00"


def test_default_graph_table_stats_profile_and_export_share_one_hundred(client, make_case):
    import io
    import json
    from openpyxl import load_workbook

    cid = make_case("Synthetic default event contract")
    observations = imported(client, cid, pair("SYN-ONE"))
    assert len(observations) == 2
    track(client, cid, observations)
    live = graph(client, cid)
    projected = projection(client, cid)
    stats = client.get(f"/api/v1/cases/{cid}/stats").json()
    legacy = client.get("/api/stats", params={"case_id": cid}).json()
    table = client.get("/api/flows", params={"case_id": cid, "tracked": 1}).json()
    v1table = client.get(f"/api/v1/cases/{cid}/flows").json()
    for value in (live, projected, stats, legacy, table, v1table):
        assert value["funds"] == live["funds"]
        assert value["analysis_revision"] == live["analysis_revision"]
    assert live["funds"]["observation_count"] == 2
    assert live["funds"]["count"] == {"lower": 1, "upper": 1}
    assert live["funds"]["totals"] == [{"currency": "CNY", "amount": "100.00", "lower": "100.00", "upper": "100.00"}]
    assert len(table["rows"]) == len(v1table["items"]) == 1
    expected_ids = {row["id"] for row in observations}
    assert set(table["rows"][0]["observation_ids"]) == set(v1table["items"][0]["observation_ids"]) == expected_ids
    selected = client.get("/api/flows", params={"case_id": cid, "ids": observations[1]["id"]}).json()["rows"]
    assert selected[0]["observation_ids"] == [observations[1]["id"]]
    assert table["rows"][0]["balance"] is None  # An event has no single own-book balance.
    sources = client.get(f"/api/v1/cases/{cid}/flows/{table['rows'][0]['id']}/sources").json()
    assert sources["total"] == 2
    assert {row["observation_balance"] for row in sources["items"]} == {"50.00", "200.00"}
    assert {row["observation_flow_id"] for row in sources["items"]} == {row["id"] for row in observations}
    payer = observations[0]["payer_account"]
    profile = client.get(f"/api/accounts/{payer}/profile", params={"case_id": cid}).json()
    assert profile["out_amt"] == "100.00" and profile["out_cnt"] == 1
    assert profile["top_counterparties"][0]["out_amt"] == "100.00"
    registry = client.get("/api/account-registry/profile", params={"case_id": cid, "account": payer}).json()
    assert registry["out_funds"] == profile["out_funds"]
    assert registry["out_cnt"] == 1 and registry["out_amt"] == "100.00"
    assert registry["top_counterparties"][0]["out_funds"] == profile["out_funds"]
    response = client.get("/api/export/tracked.xlsx", params={"case_id": cid, "expected_revision": live["analysis_revision"]})
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert book["追踪流水"].max_row == 2 and book["原始账簿观察"].max_row == 3
    snapshot = dict(book["分析修订"].values)
    assert json.loads(snapshot["统一资金投影"]) == live["funds"]
    originals = list(book["原始账簿观察"].values)
    assert [row[3] for row in originals[1:]] == ["100.00", "100.00"]
