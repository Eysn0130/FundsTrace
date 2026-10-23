"""Automatic import facts remain effective in newer main account/canvas paths."""
import io

from openpyxl import load_workbook

from test_automatic_table_import import preview
from test_feedback_import_api import complete
from test_cancelled_analysis import seed


def test_main_account_workspace_uses_effective_money_and_keeps_original_provenance(
        client, database, make_case, flow_row):
    cid, anchor, cancelled, kept, refund = seed(database, make_case, flow_row)
    registry = client.get('/api/account-registry', params={'case_id': cid, 'source': 'flows'})
    assert registry.status_code == 200
    assert {row['account'] for row in registry.json()['items']} == {'A', 'B', 'C'}
    profile = client.get('/api/account-registry/profile', params={'case_id': cid, 'account': 'B'})
    assert profile.status_code == 200
    value = profile.json()
    assert (value['cnt'], value['out_amt'], value['in_amt']) == (3, '10.00', '13.00')
    assert {row['acct'] for row in value['top_counterparties']} == {'A', 'C'}
    assert database.q_raw('select id from flows where case_id=? and id=?', [cid, cancelled]) == [{'id': cancelled}]
    sources = client.get(f'/api/v1/cases/{cid}/flows/{cancelled}/sources')
    assert sources.status_code == 200 and sources.json()['items'][0]['source_id'] == 'synthetic-v2'


def test_main_analysis_canvas_creation_keeps_its_retry_contract_and_rejects_excluded_money(
        client, database, make_case, flow_row):
    cid, anchor, cancelled, kept, refund = seed(database, make_case, flow_row)
    before = database.q_raw('select * from canvases where case_id=? order by id', [cid])
    body = {'case_id': cid, 'name': 'Synthetic effective analysis', 'flow_ids': [cancelled], 'snapshot': {}}
    assert client.post('/api/canvases/from-flows', json=body).status_code == 404
    assert database.q_raw('select * from canvases where case_id=? order by id', [cid]) == before

    body['flow_ids'] = [anchor, kept, refund]
    created = client.post('/api/canvases/from-flows', json=body)
    assert created.status_code == 200 and created.json()['added'] == 3
    repeated = client.post('/api/canvases/from-flows', json=body)
    assert repeated.status_code == 200
    assert (repeated.json()['canvas_id'], repeated.json()['added']) == (created.json()['canvas_id'], 0)
    graph = client.get('/api/graph', params={'case_id': cid, 'canvas_id': created.json()['canvas_id']}).json()
    assert set(graph['tracked_ids']) == {anchor, kept, refund}
    assert {flow for edge in graph['edges'] for flow in edge['ids']} == {anchor, kept, refund}
    database.x("update flows set payee_account='' where case_id=? and id=?", [cid, anchor])
    body.update(name='Synthetic missing counterparty', flow_ids=[anchor])
    before = database.q_raw('select * from canvases where case_id=? order by id', [cid])
    assert client.post('/api/canvases/from-flows', json=body).status_code == 422
    assert database.q_raw('select * from canvases where case_id=? order by id', [cid]) == before


def test_table_only_tracked_export_discloses_unresolved_source_coverage(
        client, database, make_case):
    cid = make_case('Synthetic table export coverage')
    data = (b'date,account,counterparty account,amount,direction\n'
            b'2026-01-01,SYN-OWN,SYN-OTHER,10.00,out\n'
            b'2026-01-02,SYN-OWN,SYN-UNKNOWN,99.00,unknown\n')
    view = preview(client, cid, data)
    assert view['automatic'] and view['table_summary']['unresolved'] == 1
    job = complete(client, cid, view)
    assert job['stage'] == 'done' and job['result']['ok'] == 1
    ids = [row['id'] for row in database.q_raw('select id from flows where case_id=?', [cid])]
    assert client.post('/api/track', json={'case_id': cid, 'flow_ids': ids}).status_code == 200
    response = client.get('/api/export/tracked.xlsx', params={'case_id': cid})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert '统计口径' in book.sheetnames
        notes = dict(list(book['统计口径'].values)[1:])
        assert notes['资金效果未确定记录数'] == 1
        assert notes['已导入来源资金语义是否完整'] == '否'
        assert '未计入流水统计' in notes['未确定记录']
    finally:
        book.close()
