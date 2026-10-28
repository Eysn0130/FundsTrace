"""Public investigation/export coverage for unresolved historical observations."""
import io
import json
from decimal import Decimal

from openpyxl import load_workbook
from test_canvas_investigation import seed, request


def test_path_steps_and_currency_scope_use_the_same_unresolved_funds(client, database, make_case):
    cid = make_case("Synthetic path evidence gaps")
    seed(database, cid, [{"amount": Decimal("100.00")},
                         {"self_account": "SYN-B", "amount": Decimal("100.00")},
                         {"currency": "USD", "amount": Decimal("3.00")}])
    value = request(client, cid, endpoint="paths", start_account="SYN-A", end_account="SYN-B", currency="CNY")
    step = value["paths"][0]["steps"][0]
    assert step["amount"] is step["count"] is None
    assert step["funds"] == value["summary"]["funds"]
    assert step["funds"]["count"] == {"lower": 0, "upper": 2}
    assert step["funds"]["totals"] == [{"currency": "CNY", "amount": None, "lower": "0.00", "upper": "200.00"}]
    window = request(client, cid)
    sender = next(item for item in window["nodes"] if item["account"] == "SYN-A")
    assert sender["currency_totals"] == [{"currency": "CNY", "in_amt": "0.00", "out_amt": None},
                                         {"currency": "USD", "in_amt": "0.00", "out_amt": None}]


def test_missing_amount_remains_in_time_scope_and_exports_as_missing(client, database, make_case):
    cid = make_case("Synthetic unknown amount")
    seed(database, cid, [{"amount": None}])
    graph = client.get("/api/graph", params={"case_id": cid}).json()
    window = request(client, cid)
    assert window["funds"] == graph["funds"]
    assert window["funds"]["observation_count"] == 1
    assert window["funds"]["totals"][0]["upper"] is None
    assert window["flows"][0]["amount"] is None
    flows = client.get(f"/api/v1/cases/{cid}/flows")
    assert flows.status_code == 200 and flows.json()["items"][0]["amount"] is None
    response = client.get("/api/export/tracked.xlsx", params={"case_id": cid})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    rows = list(book["追踪流水"].values)
    assert rows[1][rows[0].index("交易金额")] is None
    assert "nan" not in str(list(book["原始账簿观察"].values)).lower()


def test_unknown_currency_paths_do_not_mix_with_known_currency(client, database, make_case):
    cid = make_case("Synthetic unclassified currencies")
    seed(database, cid, [{"currency": None}, {"currency": "CNY"},
                         {"payer_account": "SYN-B", "payee_account": "SYN-C", "currency": "CNY"}])
    value = request(client, cid, endpoint="paths", start_account="SYN-A", end_account="SYN-C")
    assert len(value["paths"]) == 1 and value["paths"][0]["currency"] == "CNY"
    assert {step["currency"] for step in value["paths"][0]["steps"]} == {"CNY"}


def test_unclassified_direction_cannot_be_zero_in_a_known_currency(client, database, make_case):
    cid = make_case("Synthetic unknown directional currency")
    seed(database, cid, [{"currency": "CNY"},
                         {"payer_account": "SYN-B", "payee_account": "SYN-C", "currency": None}])
    for value in (client.get("/api/graph", params={"case_id": cid}).json(), request(client, cid)):
        node = next(item for item in value["nodes"] if item["account"] == "SYN-B")
        cny = next(item for item in node["currency_totals"] if item["currency"] == "CNY")
        assert cny["out_amt"] is None
    exported = client.get("/api/export/tracked.xlsx", params={"case_id": cid})
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    try:
        values = list(book["待调取账号"].values)
        pending = [dict(zip(values[0], row)) for row in values[1:]]
        cny = next(row for row in pending if row["账号"] == "SYN-B" and row["币种"] == "CNY")
        assert cny["转出金额"] == cny["交易总额"] == "≥0.00（上界未知）"
        assert cny["转出笔数"] == "0—1"
    finally:
        book.close()


def test_legacy_package_null_amount_survives_public_restore_table_time_and_export(client, database, make_case, flow_row):
    from test_cases import populated_case, exported, members, archive_of, import_package
    cid = populated_case(database, make_case, flow_row, name="Synthetic nullable legacy package")
    database.x("update flows set amount=null where case_id=?", [cid])
    entries = members(exported(client, cid))
    meta = json.loads(entries["meta.json"])
    meta["format"] = 1
    meta.pop("source_files")
    entries["meta.json"] = json.dumps(meta).encode()
    entries.pop("import_sources.parquet")
    entries.pop("flow_sources.parquet")
    restored = import_package(client, archive_of(entries))
    assert restored.status_code == 200, restored.text
    restored_id = restored.json()["case_id"]
    table = client.get(f"/api/v1/cases/{restored_id}/flows")
    assert table.status_code == 200 and all(row["amount"] is None for row in table.json()["items"])
    window = request(client, restored_id)
    graph = client.get("/api/graph", params={"case_id": restored_id}).json()
    assert window["funds"] == graph["funds"] and window["funds"]["status"] == "unknown"
    assert client.get("/api/export/tracked.xlsx", params={"case_id": restored_id}).status_code == 200


def test_original_ledger_export_keeps_raw_identifiers_after_typed_event_resolution(client, make_case):
    from test_event_import import csv_rows, transaction
    from test_folder_evidence import accepted
    cid = make_case("Synthetic original export identifiers")
    accepted(client, cid, [("books.csv", csv_rows([transaction("SYN-A"), transaction("SYN-B")]))], scope_bank=None)
    rows = client.get(f"/api/v1/cases/{cid}/flows", params={"granularity": "observations"}).json()["items"]
    assert len(rows) == 2 and all(row["self_account"].startswith("fte_") for row in rows)
    assert client.post("/api/track", json={"case_id": cid, "flow_ids": [row["id"] for row in rows]}).status_code == 200
    exported = client.get("/api/export/tracked.xlsx", params={"case_id": cid})
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    try:
        values = list(book["原始账簿观察"].values)
        original = [dict(zip(values[0], row)) for row in values[1:]]
        assert sorted((row["本方账号"], row["付款账号"], row["收款账号"], row["观察金额"], row["本方原余额"], row["原流水号"])
                      for row in original) == [
            ("SYN-A", "SYN-A", "SYN-B", "100.00", "50.00", "SYN-LOCAL-SYN-A"),
            ("SYN-B", "SYN-A", "SYN-B", "100.00", "200.00", "SYN-LOCAL-SYN-B")]
        default = list(book["追踪流水"].values)
        assert len(default) == 2
        assert default[1][default[0].index("交易金额")] == "100.00"
    finally:
        book.close()
