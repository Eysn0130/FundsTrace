"""Exact money across public legacy/v1 APIs, currency grouping and exports."""
from decimal import Decimal
import io
import json
import re
from concurrent.futures import ThreadPoolExecutor
from threading import Event

import openpyxl
import pytest

from funds_assertions import unresolved_funds
from app import importer
from app.money import money_decimal, money_string


def seed(database, make_case, flow_row):
    cid = make_case("Synthetic exact legacy money")
    amounts = ["9999999999999999.99", "9999999999999999.98", "0.01", "0.10", "0.20"]
    rows = [flow_row(trade_no=f"MONEY-{i}", amount=Decimal(amount), balance=Decimal("-0.01"),
                     trade_time=f"2026-01-01 10:{i:02d}:00") for i, amount in enumerate(amounts)]
    importer.commit_rows("synthetic-money.csv", "A", rows, case_id=cid)
    ids = [row["id"] for row in database.q("select id from flows where case_id=? order by id", [cid])]
    return cid, amounts, ids


def test_legacy_and_v1_return_identical_exact_amounts_and_filter_boundaries(client, database, make_case, flow_row):
    cid, amounts, ids = seed(database, make_case, flow_row)
    legacy = client.get("/api/flows", params={"case_id": cid}).json()["rows"]
    versioned = client.get(f"/api/v1/cases/{cid}/flows").json()["items"]
    assert {row["amount"] for row in legacy} == set(amounts)
    assert {row["amount"] for row in versioned} == set(amounts)
    assert all(row["balance"] == "-0.01" for row in legacy)
    for boundary in amounts:
        result = client.get("/api/flows", params={"case_id": cid, "min_amount": boundary, "max_amount": boundary})
        assert result.status_code == 200
        assert [row["amount"] for row in result.json()["rows"]] == [boundary]
    assert client.get("/api/flows", params={"case_id": cid, "min_amount": "9999999999999999.98"}).json()["total"] == 2
    other = make_case("Synthetic isolated money")
    assert client.get("/api/flows", params={"case_id": other, "min_amount": "0.01"}).json()["rows"] == []
    assert client.get(f"/api/context/{ids[0]}", params={"case_id": other}).status_code == 404


@pytest.mark.parametrize("query", [{"min_amount": "NaN"}, {"max_amount": "Infinity"},
    {"min_amount": "0.001"}, {"min_amount": "0.11", "max_amount": "0.10"}])
def test_invalid_money_filter_rejects_instead_of_rounding(client, make_case, query):
    assert client.get("/api/flows", params={"case_id": make_case(), **query}).status_code == 422


def test_graph_profile_context_and_excel_keep_cents_without_currency_mixing(client, database, make_case, flow_row):
    cid, amounts, ids = seed(database, make_case, flow_row)
    # Same parties/amount/time but another currency must not be graph-deduplicated.
    importer.commit_rows("synthetic-usd.csv", "A", [flow_row(trade_no="USD-1", currency="USD",
        amount=Decimal(amounts[0]), balance=None, self_account="SYN-A")], case_id=cid)
    ids = [row["id"] for row in database.q("select id from flows where case_id=? order by id", [cid])]
    assert client.post("/api/track", json={"case_id": cid, "flow_ids": ids}).status_code == 200
    graph = client.get("/api/graph", params={"case_id": cid}).json()
    edge = graph["edges"][0]
    expected = [{"currency": "CNY", "amount": None}, {"currency": "USD", "amount": None}]
    assert edge["funds"] == unresolved_funds(6, ("CNY", "20000000000000000.28"), ("USD", "9999999999999999.99"))
    assert edge["cnt"] is None
    assert edge["total"] is None
    assert edge["amount_totals"] == expected
    sender = next(row for row in graph["nodes"] if row["account"] == "SYN-A")
    assert sender["out_amt"] is None and sender["out_totals"] == expected
    profile = client.get("/api/accounts/SYN-A/profile", params={"case_id": cid}).json()
    assert profile["out_amt"] is None and profile["out_totals"] == expected
    assert profile["top_counterparties"][0]["out_totals"] == expected
    context = client.get(f"/api/context/{ids[0]}", params={"case_id": cid}).json()
    assert context["anchor"]["amount"] == amounts[0]
    assert {row["currency"] for row in context["rows"]} == {"CNY", "USD"}
    workbook = openpyxl.load_workbook(io.BytesIO(client.get("/api/export/tracked.xlsx", params={"case_id": cid}).content))
    sheet = workbook["追踪流水"]
    header = [cell.value for cell in sheet[1]]
    money_column = header.index("交易金额") + 1
    cells = [sheet.cell(row, money_column) for row in range(2, sheet.max_row + 1)]
    assert all(cell.data_type == "s" for cell in cells)
    assert [cell.value for cell in cells] == [amounts[0], amounts[0], *amounts[1:]]
    pending = list(workbook["待调取账号"].values)
    headers = pending[0]
    records = [dict(zip(headers, row)) for row in pending[1:]]
    assert {row["币种"] for row in records} == {"CNY", "USD"}
    assert all(isinstance(row["交易总额"], str) for row in records)
    # Self-contained graph snapshot keeps both authoritative strings and labels.
    payload = {"case_name": "Synthetic money", "nodes": [], "edges": [{"source": "A", "target": "B",
        "label": "CNY 20,000,000,000,000,000.28", "data": edge}]}
    exported = client.post("/api/export/graph.html", json=payload)
    encoded = re.search(r"var DATA = (.*?);\nfunction buildWM", exported.text, re.S)
    assert json.loads(encoded.group(1))["edges"][0]["data"]["amount_totals"] == expected


@pytest.mark.parametrize("value,expected", [("1e3", "1000.00"), ("1.23E+5", "123000.00"),
    ("￥1,234.565", "1234.57"), ("(0.01)", "-0.01"), ("0.005", "0.01")])
def test_import_decimal_notation_is_parsed_whole(value, expected):
    assert importer.parse_amount(value) == Decimal(expected)


@pytest.mark.parametrize("value", ["12junk34", "1.2.3", "1e9999", "NaN", "Infinity"])
def test_import_rejects_corrupted_money(value):
    assert importer.parse_amount(value) is None


def test_money_boundary_never_accepts_binary_float():
    with pytest.raises(ValueError):
        money_decimal(0.1)
    assert money_string(Decimal("-0.00")) == "0.00"


def test_unknown_currencies_remain_unbounded_without_fabricated_currency_group(client, database, make_case):
    cid = make_case("Synthetic unknown currency")
    for currency, amount in [(None, "1.01"), ("", "2.02")]:
        database.x("insert into flows(case_id,trade_time,payer_account,payee_account,amount,currency) "
                   "values (?,'2026-01-01 10:00:00','UNKNOWN-A','UNKNOWN-B',?,?)", [cid, Decimal(amount), currency])
    ids = [row["id"] for row in database.q("select id from flows where case_id=?", [cid])]
    client.post("/api/track", json={"case_id": cid, "flow_ids": ids})
    edge = client.get("/api/graph", params={"case_id": cid}).json()["edges"][0]
    assert edge["total"] is None and edge["amount_totals"] == [{"currency": None, "amount": None}]
    assert edge["funds"] == unresolved_funds(2, (None, None))
    profile = client.get("/api/accounts/UNKNOWN-A/profile", params={"case_id": cid}).json()
    assert profile["out_totals"] == [{"currency": None, "amount": None}]
    assert profile["out_funds"] == edge["funds"]
    originals = client.get("/api/flows", params={"case_id": cid, "granularity": "observations"}).json()["rows"]
    assert sorted((row["id"], row["amount"], row["currency"]) for row in originals) == [
        (ids[0], "1.01", None), (ids[1], "2.02", "")]
    workbook = openpyxl.load_workbook(io.BytesIO(client.get("/api/export/tracked.xlsx", params={"case_id": cid}).content))
    records = list(workbook["待调取账号"].values)
    assert len(records) == 3  # One row per account, rather than duplicated NULL/empty groups.


def test_case_package_round_trip_keeps_exact_money(client, database, make_case, flow_row):
    cid, amounts, _ids = seed(database, make_case, flow_row)
    package = client.get(f"/api/cases/{cid}/export")
    restored = client.post("/api/cases/import", files={"file": ("synthetic-money.fundcase", package.content)})
    assert restored.status_code == 200
    new_id = restored.json()["case_id"]
    assert new_id != cid
    rows = client.get("/api/flows", params={"case_id": new_id}).json()["rows"]
    assert {row["amount"] for row in rows} == set(amounts)
    assert all(row["balance"] == "-0.01" for row in rows)


def test_context_without_transaction_time_has_no_timed_suggestions(client, database, make_case):
    cid = make_case("Synthetic missing time")
    fid = database.q("insert into flows(case_id,payer_account,payee_account,amount) "
                     "values (?,'SYN-A','SYN-B',0.01) returning id", [cid])[0]["id"]
    result = client.get(f"/api/context/{fid}", params={"case_id": cid})
    assert result.status_code == 200
    assert result.json()["anchor"]["amount"] == "0.01"
    assert result.json()["suggestions"] == {"sources": [], "targets": []}


def test_graph_reads_one_state_while_background_delete_waits(client, database, make_case, flow_row, monkeypatch):
    cid = make_case("Synthetic concurrent graph")
    importer.commit_rows("synthetic-lock.csv", "A", [flow_row(amount=Decimal("0.01"))], case_id=cid)
    fid = database.q("select id from flows where case_id=?", [cid])[0]["id"]
    client.post("/api/track", json={"case_id": cid, "flow_ids": [fid]})
    read = database.q_raw
    attempted, finished = Event(), Event()
    futures = []
    with ThreadPoolExecutor(max_workers=1) as pool:
        def delete():
            attempted.set()
            with database.transaction():
                database.x("delete from tracked_flows where case_id=?", [cid])
                database.x("delete from flows where case_id=?", [cid])
            finished.set()

        def interleave(sql, params=None):
            rows = read(sql, params)
            if "from selected_facts" in sql and not futures:
                futures.append(pool.submit(delete))
                assert attempted.wait(1)
                finished.wait(0.05)  # Without the whole-read lock deletion finishes before the next query.
            return rows

        monkeypatch.setattr(database, "q_raw", interleave)
        result = client.get("/api/graph", params={"case_id": cid}).json()
        assert result["tracked_ids"] == [fid]
        assert result["edges"][0]["cnt"] is None
        assert result["edges"][0]["total"] is None
        assert result["edges"][0]["funds"] == unresolved_funds(1, ("CNY", "0.01"))
        assert result["edges"][0]["ids"] == [fid]
        futures[0].result(timeout=2)
    assert finished.is_set()
