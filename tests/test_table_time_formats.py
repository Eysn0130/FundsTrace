"""Synthetic public-parser expectations for declared time formats."""
import csv
import datetime as dt
import io

from openpyxl import Workbook
import pytest

from app import table_import as tables, table_time_formats as formats


DECLARATIONS = (
    ("YYYY-MM-DD", "2026-01-02", dt.datetime(2026, 1, 2)),
    ("YYYY/MM/DD", "2026/01/02", dt.datetime(2026, 1, 2)),
    ("YYYYMMDD", "20260102", dt.datetime(2026, 1, 2)),
    ("YYYY-MM-DD HH:MM:SS", "2026-01-02 12:34:56", dt.datetime(2026, 1, 2, 12, 34, 56)),
    ("YYYYMMDDHHMMSS", "20260102123456", dt.datetime(2026, 1, 2, 12, 34, 56)),
    ("YYYYMMDD HHMMSS", "20260102 123456", dt.datetime(2026, 1, 2, 12, 34, 56)),
)


def plan(headers, records, *, suffix=".csv", **options):
    if suffix == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        for row in (headers, *records):
            sheet.append(row)
        output = io.BytesIO()
        workbook.save(output)
        payload = output.getvalue()
    else:
        output = io.StringIO(newline="")
        csv.writer(output).writerows([headers, *records])
        payload = output.getvalue().encode()
    options.setdefault("profile", tables.TIME_FORMAT_PROFILE)
    return tables.parse(payload, suffix=suffix, **options)


def headers(time_header):
    return [time_header, "本方账号", "对方账号", "收支标志", "交易金额"]


def row(time):
    return [time, "SYN-OWN", "SYN-COUNTER", "收入", "1.23"]


@pytest.mark.parametrize("spec,value,expected", DECLARATIONS)
@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_declared_date_and_datetime_headers_automatically_import(spec, value, expected, suffix):
    result = plan(headers(f"交易时间（{spec}）"), [row(value)], suffix=suffix)
    assert result["profile"] == "funds-trace-table-v10"
    assert result["stats"]["accepted"] == 1
    assert result["stats"]["unresolved"] == 0
    assert result["rows"][0]["trade_time"] == expected


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
@pytest.mark.parametrize("spec,value", [("HH:MM:SS", "12:34:56"), ("HHMMSS", "123456")])
def test_declared_split_date_and_clock_combine_without_guesses(spec, value, suffix):
    result = plan(["交易日期(YYYYMMDD)", f"交易时分秒[{spec}]", *headers("")[1:]],
                  [["20260102", value, *row("")[1:]]], suffix=suffix)
    assert result["stats"]["accepted"] == 1
    assert result["rows"][0]["trade_time"] == dt.datetime(2026, 1, 2, 12, 34, 56)


@pytest.mark.parametrize("profile", [f"funds-trace-table-v{version}" for version in range(5, 10)])
@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_historical_profiles_do_not_gain_new_header_meaning(profile, suffix):
    result = plan(headers("交易日期(YYYY-MM-DD)"), [row("2026-01-02")], suffix=suffix, profile=profile)
    assert result["profile"] == profile
    assert not result["rows"]
    assert result["stats"]["reason_counts"] == {"unknown_table": 2}


def test_all_rows_validate_against_declared_format_after_first_hundred():
    result = plan(headers("交易日期(YYYY-MM-DD)"),
                  [row("2026-01-02") for _ in range(105)] + [row("20260103")])
    assert result["stats"]["accepted"] == 105
    assert result["stats"]["unresolved"] == 1
    assert result["records"][-1]["reason_code"] == "invalid_time"


def test_new_declared_header_rebinds_following_region():
    result = plan(headers("交易日期(YYYY-MM-DD)"), [
        row("2026-01-02"), headers("交易日期[YYYYMMDD]"), row("20260103"), row("2026-01-04"),
    ])
    assert [record["reason_code"] for record in result["records"]] == [
        "transaction_valid", "repeated_header", "transaction_valid", "invalid_time",
    ]
    assert [value["trade_time"] for value in result["rows"]] == [dt.datetime(2026, 1, 2), dt.datetime(2026, 1, 3)]


def test_conflicting_full_candidate_is_not_ignored_by_annotation():
    result = plan(["交易日期(YYYY-MM-DD)", "交易时间(YYYYMMDDHHMMSS)", *headers("")[1:]],
                  [["2026-01-02", "20260103123456", *row("")[1:]]])
    assert not result["rows"]
    assert result["records"][0]["reason_code"] == "conflicting_time"


@pytest.mark.parametrize("header", [
    "交易日期(DD/MM/YYYY)", "交易日期(任意格式)", "交易日期(YYYY-MM-DD)备注",
    "交易日期(YYYY-MM-DD](", "未知字段(YYYY-MM-DD)",
])
def test_unknown_or_malformed_declarations_do_not_become_headers(header):
    result = plan(headers(header), [row("2026-01-02")])
    assert not result["rows"]
    assert result["stats"]["reason_counts"] == {"unknown_table": 2}


@pytest.mark.parametrize("header,value", [
    ("交易日期(HHMMSS)", "123456"),
    ("交易日期[HH:MM:SS]", "12:34:56"),
    ("交易时分秒(YYYYMMDD)", "20260102"),
    ("交易时分秒(YYYY-MM-DD HH:MM:SS)", "2026-01-02 12:34:56"),
])
def test_format_must_match_the_explicit_header_role(header, value):
    result = plan(headers(header), [row(value)])
    assert not result["rows"]
    assert result["stats"]["reason_counts"] == {"unknown_table": 2}


def test_declared_clock_validates_every_row_including_last():
    result = plan(["交易日期(YYYYMMDD)", "交易时间(HHMMSS)", *headers("")[1:]],
                  [["20260102", "123456", *row("")[1:]] for _ in range(105)]
                  + [["20260102", "12:34:56", *row("")[1:]]])
    assert result["stats"]["accepted"] == 105
    assert result["stats"]["unresolved"] == 1
    assert result["records"][-1]["reason_code"] == "invalid_time"


@pytest.mark.parametrize("header,expected", [
    ("交易日期（YYYY-MM-DD）", ("交易日期", "YYYY-MM-DD")),
    ("transaction time [yyyy-mm-dd hh:mm:ss]", ("transaction time", "YYYY-MM-DD HH:MM:SS")),
    ("交易时分秒(HHMMSS)", ("交易时分秒", "HHMMSS")),
    ("交易日期(格式:YYYY-MM-DD)", None),
    ("交易日期(YYYY-MM-DD)备注", None),
    ("交易日期(YYYY-MM-DD]", None),
    ("交易日期(YYYY-MM-DD)(HHMMSS)", None),
    ("(YYYY-MM-DD)", None),
    (None, None),
])
def test_header_grammar_is_closed_and_keeps_unknown_base_for_caller(header, expected):
    assert formats.header_parts(header) == expected


@pytest.mark.parametrize("spec,value,_", DECLARATIONS)
def test_declared_lexical_formats_validate(spec, value, _):
    assert formats.validate(value, spec)


@pytest.mark.parametrize("value,spec", [
    ("2026-1-02", "YYYY-MM-DD"), ("2026-02-29", "YYYY-MM-DD"),
    ("2026-01-02", "YYYYMMDD"), ("２０２６-０１-０２", "YYYY-MM-DD"),
    ("20260102123456.000001", "YYYYMMDDHHMMSS"),
    ("2026-01-02 12:34:56.000001", "YYYY-MM-DD HH:MM:SS"),
    ("2026-01-02T12:34:56Z", "YYYY-MM-DD HH:MM:SS"),
    ("2026-01-02 24:00:00", "YYYY-MM-DD HH:MM:SS"),
    ("240000", "HHMMSS"), ("12345", "HHMMSS"),
    ("12:34", "HH:MM:SS"), ("12:34:56.1", "HH:MM:SS"),
    (46024, "YYYYMMDD"), (0.5, "HHMMSS"), (True, "YYYYMMDD"),
    ("2026-01-02", "unsupported"),
])
def test_declaration_never_repairs_or_discards_information(value, spec):
    assert not formats.validate(value, spec)


@pytest.mark.parametrize("value,spec,date_only,expected", [
    (dt.date(2026, 1, 2), "YYYY-MM-DD", False, True),
    (dt.datetime(2026, 1, 2), "YYYYMMDD", True, True),
    (dt.datetime(2026, 1, 2, 1), "YYYYMMDD", False, False),
    (dt.datetime(2026, 1, 2, 1), "YYYY-MM-DD HH:MM:SS", False, True),
    (dt.datetime(2026, 1, 2), "YYYY-MM-DD HH:MM:SS", True, False),
    (dt.date(2026, 1, 2), "YYYY-MM-DD HH:MM:SS", False, False),
    (dt.datetime(2026, 1, 2, 1, microsecond=1), "YYYY-MM-DD HH:MM:SS", False, False),
    (dt.datetime(2026, 1, 2, tzinfo=dt.timezone.utc), "YYYY-MM-DD HH:MM:SS", False, False),
    (dt.time(12, 34, 56), "HHMMSS", False, True),
    (dt.time(12, 34, 56), "HHMMSS", True, False),
    (dt.time(12, 34, 56, 1), "HHMMSS", False, False),
    (dt.time(12, 34, 56, tzinfo=dt.timezone.utc), "HHMMSS", False, False),
    (dt.datetime(2026, 1, 2, 12, 34, 56), "HHMMSS", False, False),
])
def test_typed_excel_semantics_are_not_coerced(value, spec, date_only, expected):
    assert formats.validate(value, spec, date_only=date_only) is expected


@pytest.mark.parametrize("value,number_format,expected", [
    (dt.date(2026, 1, 2), "yyyy-mm-dd", "transaction_valid"),
    (dt.datetime(2026, 1, 2), "yyyy-mm-dd", "transaction_valid"),
    (dt.datetime(2026, 1, 2, 12), "yyyy-mm-dd hh:mm:ss", "invalid_time"),
])
def test_typed_excel_date_declaration_preserves_clock_precision(value, number_format, expected):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers("交易日期(YYYY-MM-DD)"))
    sheet.append(row(value))
    sheet.cell(2, 1).number_format = number_format
    output = io.BytesIO()
    workbook.save(output)
    result = tables.parse(output.getvalue(), suffix=".xlsx", profile=tables.TIME_FORMAT_PROFILE)
    assert result["records"][0]["reason_code"] == expected


@pytest.mark.parametrize("date_value,expected", [
    (dt.datetime(2026, 1, 2), "transaction_valid"),
    (dt.datetime(2026, 1, 2, 1), "invalid_time"),
])
def test_explicit_typed_date_can_combine_clock_without_erasing_nonzero_time(date_value, expected):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["交易日期(YYYY-MM-DD)", "交易时分秒(HHMMSS)", *headers("")[1:]])
    sheet.append([date_value, dt.time(12, 34, 56), *row("")[1:]])
    sheet.cell(2, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    sheet.cell(2, 2).number_format = "hh:mm:ss"
    output = io.BytesIO()
    workbook.save(output)
    result = tables.parse(output.getvalue(), suffix=".xlsx", profile=tables.TIME_FORMAT_PROFILE)
    assert result["records"][0]["reason_code"] == expected
    if expected == "transaction_valid":
        assert result["rows"][0]["trade_time"] == dt.datetime(2026, 1, 2, 12, 34, 56)
