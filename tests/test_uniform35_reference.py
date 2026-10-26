"""Exact business-only form of uniform-37, including stale XLSX dimensions."""
import io
import re
import zipfile

import pytest

from app import table_import
from test_feedback31_reference import payload
from test_folder_evidence import accepted
from test_uniform37_reference import COLUMNS as PADDED_COLUMNS, row


COLUMNS = PADDED_COLUMNS[:-2]


def workbook(records, headers=COLUMNS, stale=False, padding=None):
    raw = payload(records, headers, suffix="xlsx")
    if padding is not None:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw))
        for index, value in enumerate(padding, len(headers) + 1):
            book.active.cell(row=2, column=index, value=value)
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        raw = stream.getvalue()
    if not stale:
        return raw
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', content)
            target.writestr(member, content)
    return output.getvalue()


@pytest.mark.parametrize("stale", [False, True])
@pytest.mark.parametrize("padding", [None, ("", "")])
def test_whole35_workbook_preserves_references_and_reads_physical_rows(client, make_case, stale, padding):
    cid = make_case("Synthetic business-only workbook")
    raw = workbook([row(), row(**{"交易金额": "0", "交易流水号": "SYN-ZERO"})], stale=stale, padding=padding)
    view, receipt = accepted(client, cid, [("rows.xlsx", raw)], scope_bank=None)
    assert view["stats"]["ok"] == receipt["ok"] == 1
    assert receipt["table_summary"]["non_monetary"] == 1
    flow = client.get(f"/api/v1/cases/{cid}/flows").json()["items"][0]
    assert flow["trade_no"] == "SYN-PRIMARY" and flow["amount"] == "12.34" and flow["raw_self_account"] == "SYN-A"
    original = client.get(f"/api/v1/cases/{cid}/sources/{receipt['source_id']}/file")
    with zipfile.ZipFile(io.BytesIO(original.content)) as source:
        assert source.read("rows.xlsx") == raw
    package = client.get(f"/api/cases/{cid}/export").content
    restored = client.post("/api/cases/import", files={"file": ("synthetic.fundcase", package)})
    assert restored.status_code == 200, restored.text
    assert client.get(f"/api/v1/cases/{restored.json()['case_id']}/flows").json()["items"][0]["trade_no"] == "SYN-PRIMARY"


@pytest.mark.parametrize("headers", [COLUMNS[:-1], COLUMNS + ("extra",),
    tuple("交易流水号" if name == "备注" else name for name in COLUMNS), COLUMNS + ("Unnamed: 35",)])
def test_only_the_exact35_signature_inherits_reference_roles(client, make_case, headers):
    _, result = accepted(client, make_case("Synthetic non-contract workbook"), [("rows.xlsx", workbook([row()], headers, True))], scope_bank=None)
    assert result["ok"] == 0 and result["table_summary"]["unresolved"] == 1


@pytest.mark.parametrize("changes", [{"交易金额": "-12.34"}, {"交易金额": "12.345"}, {"币种": "unknown"},
    {"借贷标志": "unknown"}, {"交易时间": "unknown"}, {"交易是否成功": "unknown"},
    {"本方账号": "", "本方卡号": "", "查询账号": "SYN-QUERY-ONLY"}])
def test_35_reference_roles_do_not_relax_core_values_or_promote_queries(client, make_case, changes):
    _, result = accepted(client, make_case("Synthetic workbook admission controls"), [("rows.xlsx", workbook([row(**changes)], stale=True))], scope_bank=None)
    assert result["ok"] == 0 and result["table_summary"]["unresolved"] == 1


def test_legacy_profile_does_not_inherit35_reference_roles():
    plan = table_import.parse(workbook([row()], stale=True), suffix=".xlsx", profile=table_import.CURRENCY_REGION_PROFILE,
                              reference_policy=table_import.TRANSACTION_REFERENCE_POLICY)
    assert not plan["rows"] and plan["records"][0]["reason_code"] == "conflicting_field"


@pytest.mark.parametrize("padding", [("SYN-NONEMPTY", ""), ("", "SYN-NONEMPTY"), (0, ""), ("", "=1+1"), ("",), ("", "", "")])
def test_only_two_empty_trailing_workbook_coordinates_inherit35_roles(client, make_case, padding):
    _, result = accepted(client, make_case("Synthetic worksheet padding controls"),
                         [("rows.xlsx", workbook([row()], stale=True, padding=padding))], scope_bank=None)
    assert result["ok"] == 0 and result["table_summary"]["unresolved"] == 1
