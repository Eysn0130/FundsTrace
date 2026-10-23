"""Independent synthetic expectations for the versioned full-table interpreter."""
import csv
import datetime as dt
from decimal import Decimal
import io
import re
import zipfile

from openpyxl import Workbook
import pytest

from app import table_import as tables, table_contracts


def csv_plan(headers, records, *, encoding="utf-8", delimiter=",", prefix=(), **options):
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=delimiter)
    writer.writerows([*prefix, headers, *records])
    return tables.parse(output.getvalue().encode(encoding), suffix=".csv", **options)


def a(headers=None, records=None, **options):
    return csv_plan(headers or ["date", "payer_account", "payee_account", "amount"],
                    records if records is not None else [["2026-01-01 10:11:12.123456", "00012345678901234567890", "SYN-B", "12.34"]], **options)


def b(amount="12.34", direction="收入", headers=None, **options):
    return csv_plan(headers or ["交易时间", "本方账号", "对方账号", "收支标志", "交易金额"],
                    [["2026-01-01 10:00:00", "SYN-OWN", "SYN-COUNTER", direction, amount]], **options)


def assert_conserved(plan):
    assert len(plan["records"]) == sum(member["row_count"] for member in plan["members"])
    assert plan["stats"]["total"] == sum(plan["stats"][key] for key in ("accepted", "unresolved", "non_monetary", "non_transaction", "auxiliary"))
    detail_fields = tables.DETAILS_FIELDS | (table_contracts.DETAILS_FIELDS
                                            if plan["profile"] in tables.CONTRACT_PROFILES else frozenset())
    assert all(record["details"].keys() <= detail_fields for record in plan["records"])
    if plan["profile"] in tables.CONTRACT_PROFILES:
        assert all(table_contracts.validate_details(record["details"], record["disposition"], record["reason_code"])
                   for record in plan["records"])
    assert all(isinstance(value, str) for record in plan["records"] for value in record["details"].values())
    assert all(record["reason_code"] in tables.REASON_CODES for record in plan["records"])
    assert [record["record_no"] for record in plan["records"]] == list(range(1, len(plan["records"]) + 1))
    assert all((record["disposition"] == "accepted") == ("flow_sha256" in record["details"]) for record in plan["records"])


def test_english_headers_preserve_text_and_microseconds_and_legacy_currency_basis():
    plan = a()
    assert plan["mode"] == "A"
    row = plan["rows"][0]
    assert row["amount"] == Decimal("12.34")
    assert row["payer_account"] == "00012345678901234567890"
    assert row["trade_time"] == dt.datetime(2026, 1, 1, 10, 11, 12, 123456)
    assert row["currency"] == "CNY"
    assert "历史默认" in plan["records"][0]["details"]["currency_basis"]
    assert plan["records"][0]["details"]["flow_sha256"] == tables.flow_digest(row)
    assert_conserved(plan)


@pytest.mark.parametrize("header,value,expected", [
    ("交易金额(分)", "123", "1.23"), ("交易金额（角）", "12.3", "1.23"),
    ("交易金额(万元)", "0.000123", "1.23"), ("交易金额(千元)", "0.00123", "1.23"),
    ("amount_cents", "123", "1.23"), ("amount (yuan)", "1.23", "1.23"),
    ("交易金额万元", "0.000123", "1.23"), ("交易金额(元)", "1.2300", "1.23"),
])
def test_units_convert_before_cent_precision_check(header, value, expected):
    plan = a(["交易时间", "付款账号", "收款账号", header], [["2026-01-01", "SYN-A", "SYN-B", value]])
    assert plan["rows"][0]["amount"] == Decimal(expected)
    assert_conserved(plan)


@pytest.mark.parametrize("header,value,reason", [
    ("amount", "1.234", "amount_precision"), ("交易金额(分)", "1.1", "amount_precision"),
    ("交易金额(万元)", "1000000000000", "amount_overflow"),
    ("amount", "10000000000000000", "amount_overflow"),
    ("amount", "NaN", "missing_amount"), ("amount", "abc12", "invalid_amount"),
    ("amount", "12,34", "invalid_amount"), ("交易金额(亿)", "1", "unknown_unit"),
    ("amount", "1e100000", "amount_overflow"),
    ("amount", "1e10000000", "amount_overflow"),
    ("amount", "1e-10000000", "amount_precision"),
    ("amount", "(-1)", "invalid_amount"),
])
def test_unsafe_money_is_not_rounded_or_dropped(header, value, reason):
    plan = a(["date", "payer_account", "payee_account", header], [["2026-01-01", "SYN-A", "SYN-B", value]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == reason
    assert_conserved(plan)


def test_two_amount_candidates_compare_exactly_after_units_on_every_row():
    plan = a(["date", "payer_account", "payee_account", "交易金额", "发生金额(分)"], [
        ["2026-01-01", "A", "B", "1.23", "123"],
        ["2026-01-02", "A", "B", "1.23", "124"],
    ])
    assert len(plan["rows"]) == 1
    assert plan["records"][1]["reason_code"] == "conflicting_amount"
    assert "amount" not in plan["mapping"]
    assert [evidence["candidate_count"] for evidence in plan["mapping_evidence"] if evidence["field"] == "amount"] == [2, 2]
    assert_conserved(plan)


def test_duplicate_column_positions_are_not_pandas_mangled_or_first_wins():
    plan = a(["date", "payer_account", "payee_account", "amount", "amount"], [
        ["2026-01-01", "A", "B", "10", "10"], ["2026-01-01", "A", "B", "10", "20"],
    ])
    assert plan["columns"][-2:] == ["amount [列4]", "amount [列5]"]
    assert len(plan["rows"]) == 1
    assert plan["records"][1]["reason_code"] == "conflicting_amount"


def test_column_reordering_and_synonyms_produce_same_facts():
    first = a(["date", "payer_account", "payee_account", "amount"], [["2026-01-01", "A", "B", "1.23"]])
    second = a(["收款账号", "金额(分)", "付款账号", "交易日期"], [["B", "123", "A", "2026-01-01"]])
    assert tables.flow_digest(first["rows"][0]) == tables.flow_digest(second["rows"][0])
    assert first["records"][0]["record_key"] != second["records"][0]["record_key"]


@pytest.mark.parametrize("direction,expected", [("income", "in"), ("贷", "in"), ("收入", "in"), ("D", "out"), ("支出", "out"), ("debit", "out")])
def test_explicit_direction_enum(direction, expected):
    plan = b(direction=direction)
    row = plan["rows"][0]
    assert row["payer_account"] == ("SYN-COUNTER" if expected == "in" else "SYN-OWN")
    assert row["payee_account"] == ("SYN-OWN" if expected == "in" else "SYN-COUNTER")


@pytest.mark.parametrize("direction,reason", [("", "missing_direction"), ("02", "invalid_direction"), ("不是收入", "invalid_direction"), ("支出失败", "invalid_direction")])
def test_direction_uses_full_enum_not_substring_or_default_out(direction, reason):
    plan = b(direction=direction)
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == reason


def test_negative_with_direction_is_not_blindly_abs():
    plan = b(amount="-12.34", direction="支出")
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "negative_amount"


def test_signed_header_is_explicit_convention_and_checks_independent_direction():
    plan = csv_plan(["date", "account", "signed amount"], [["2026-01-01", "A", "-12.34"], ["2026-01-02", "A", "10"]])
    assert [row["amount"] for row in plan["rows"]] == [Decimal("12.34"), Decimal("10.00")]
    assert plan["rows"][0]["payer_account"] == "A"
    assert plan["rows"][0]["payee_account"] == ""
    assert plan["rows"][1]["payee_account"] == "A"
    conflict = b(amount="-12.34", direction="收入", headers=["date", "account", "counterparty account", "direction", "signed amount"])
    assert conflict["records"][0]["reason_code"] == "conflicting_direction"
    assert not conflict["rows"]


@pytest.mark.parametrize("income,expense,direction,amount,reason", [
    ("10", "0", "收入", "10", "transaction_valid"),
    ("0", "10", "", "10", "transaction_valid"),
    ("10", "20", "收入", "10", "conflicting_income_expense"),
    ("10", "0", "支出", "10", "conflicting_direction"),
    ("10", "0", "收入", "20", "conflicting_amount"),
    ("0", "0", "", "0", "zero_amount"),
])
def test_split_income_expense_and_independent_fields_must_agree(income, expense, direction, amount, reason):
    plan = csv_plan(["date", "account", "income", "expense", "direction", "amount"], [["2026-01-01", "A", income, expense, direction, amount]])
    assert plan["records"][0]["reason_code"] == reason
    assert bool(plan["rows"]) == (reason == "transaction_valid")
    assert_conserved(plan)


def test_validation_is_full_not_first_thirty_samples_and_examples_are_only_display_limit():
    rows = [["2026-01-01", "A", "B", "10"] for _ in range(35)]
    rows += [["01/02/2026", "A", "B", "10"] for _ in range(5)]
    plan = a(records=rows)
    assert plan["stats"]["ok"] == 35
    assert plan["stats"]["unresolved"] == 5
    assert len(plan["stats"]["bad_examples"]) == 3
    assert [record["source_row"] for record in plan["records"][-5:]] == list(range(37, 42))
    assert_conserved(plan)


@pytest.mark.parametrize("encoding,delimiter", [("utf-8-sig", ","), ("gb18030", "\t"), ("utf-16", ";"), ("utf-8", "|")])
def test_encodings_delimiters_preamble_blank_repeated_header_and_bad_width(encoding, delimiter):
    header = ["交易时间", "付款账号", "收款账号", "金额"]
    plan = csv_plan(header, [["2026-01-01", "A", "B", "1"], [], header, ["2026-01-02", "A", "B", "2", "extra"], ["2026-01-03", "A"]],
                    prefix=[["测试说明"], []], encoding=encoding, delimiter=delimiter)
    assert [record["reason_code"] for record in plan["records"]] == ["preamble", "blank_record", "transaction_valid", "blank_record", "repeated_header", "row_width", "row_width"]
    assert plan["members"][0]["row_count"] == 7
    assert_conserved(plan)


def test_multiline_csv_source_rows_are_logical_records():
    plan = a(["date", "payer_account", "payee_account", "amount", "memo"], [["2026-01-01", "A", "B", "1", "line one\nline two"], ["2026-01-02", "A", "B", "2", "end"]])
    assert [record["source_row"] for record in plan["records"]] == [2, 3]


def workbook_plan(*sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return tables.parse(output.getvalue(), suffix=".xlsx")


def test_all_worksheets_have_disjoint_identities_and_unknown_sheet_is_not_lost():
    header = ["date", "payer_account", "payee_account", "amount"]
    plan = workbook_plan(("First", [header, [dt.datetime(2026, 1, 1, 12, 0, 0, 123456), "000001", "B", "1.23"]]),
                         ("Second", [header, ["2026-01-02", "A", "B", "2.34"]]),
                         ("Notes", [["unrecognized headings", "value"], ["unknown", "row"]]))
    assert len(plan["rows"]) == 2
    assert len(plan["records"]) == 4
    assert plan["sheet"] == "workbook"
    assert [row["_source_row"] for row in plan["rows"]] == [1, 2]
    assert plan["records"][0]["source_row"] == plan["records"][1]["source_row"] == 2
    assert plan["records"][0]["record_key"] != plan["records"][1]["record_key"]
    assert [member["row_count"] for member in plan["members"]] == [1, 1, 2]
    assert plan["records"][2]["disposition"] == "unresolved"
    assert plan["records"][0]["details"]["sheet"] == "First"
    assert_conserved(plan)


def test_numeric_excel_account_with_more_than_fifteen_digits_is_not_asserted_lossless():
    plan = workbook_plan(("Synthetic", [["date", "payer_account", "payee_account", "amount"], ["2026-01-01", 1234567890123456789, "B", "1.23"]]))
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "unsafe_account_number"


def test_all_unresolved_is_a_complete_plan_not_manual_mapping_error():
    plan = csv_plan(["x", "y"], [["a", "b"], ["c", "d"]])
    assert plan["stats"]["ok"] == 0
    assert plan["stats"]["unresolved"] == 3
    assert_conserved(plan)


def test_multicurrency_and_cash_indicator_is_not_currency():
    plan = a(["date", "payer_account", "payee_account", "amount", "currency", "钞汇标志"], [["2026-01-01", "A", "B", "1.23", "USD", "现钞"], ["2026-01-01", "A", "B", "2.34", "人民币", "现汇"]])
    assert plan["stats"]["currency_totals"] == {"CNY": "2.34", "USD": "1.23"}
    assert_conserved(plan)


def test_missing_counterparty_preserves_blank_account_and_does_not_forge_from_name():
    plan = csv_plan(["date", "account", "counterparty name", "direction", "amount"], [["2026-01-01", "A", "Synthetic other", "收入", "1.23"]])
    assert plan["rows"][0]["payer_account"] == ""
    assert plan["rows"][0]["payer_name"] == "Synthetic other"


def test_strict_split_date_time_and_conflicts():
    plan = csv_plan(["date", "time", "payer_account", "payee_account", "amount"], [["2026-01-01", "12:34:56.123456", "A", "B", "1.23"], ["2026-01-02 10:00:00", "11:00:00", "A", "B", "1.23"]])
    assert plan["rows"][0]["trade_time"] == dt.datetime(2026, 1, 1, 12, 34, 56, 123456)
    assert plan["records"][1]["reason_code"] == "conflicting_time"


def test_feedback_status_and_card_fallback_have_closed_meanings():
    headers = ["交易时间", "交易账号", "交易卡号", "交易户名", "交易金额", "交易余额", "收付标志", "交易是否成功", "摘要说明", "备注", "查询反馈结果原因"]
    plan = csv_plan(headers, [["2026-01-01", "", "CARD-A", "Synthetic own", "1.23", "9.87", "收", "01", "摘要甲", "备注乙", ""],
                              ["2026-01-02", "A", "CARD-A", "Synthetic own", "1", "9", "收", "02", "", "", ""],
                              ["", "A", "CARD-A", "Synthetic own", "", "", "", "", "", "", "无反馈明细"]])
    assert plan["rows"][0]["self_account"] == "CARD-A"
    assert plan["rows"][0]["self_name"] == "Synthetic own"
    assert plan["rows"][0]["summary"] == "摘要甲；备注乙"
    assert plan["records"][1]["reason_code"] == "unknown_status"
    assert plan["records"][2]["reason_code"] == "feedback_only"
    assert_conserved(plan)


def test_numeric_status_codes_are_not_assumed_for_an_unknown_generic_status_column():
    plan = a(["date", "payer_account", "payee_account", "amount", "status"], [["2026-01-01", "A", "B", "10", "1"]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "unknown_status"


def test_zero_never_becomes_missing_amount():
    plan = b(amount="0", direction="")
    assert plan["records"][0]["disposition"] == "non_monetary"
    assert plan["records"][0]["reason_code"] == "zero_amount"


@pytest.mark.parametrize("payload,suffix", [(b'"broken', ".csv"), (b"not a workbook", ".xlsx"), (b"anything", ".xls"), (b"\x00binary", ".csv")])
def test_unreadable_file_uses_safe_constant_error(payload, suffix):
    with pytest.raises(tables.TableImportError) as caught:
        tables.parse(payload, suffix=suffix)
    assert "broken" not in str(caught.value)
    assert "anything" not in str(caught.value)


def test_manual_compatibility_mapping_cannot_suppress_conflicting_amount_alias():
    plan = a(["date", "payer_account", "payee_account", "amount", "交易金额"], [["2026-01-01", "A", "B", "10", "20"]], mapping={"amount": "amount"})
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "conflicting_amount"


@pytest.mark.parametrize("value", ["-1", "0 1", "+1", "01-", "suc cess"])
def test_status_values_do_not_use_header_normalization(value):
    plan = a(["date", "payer_account", "payee_account", "amount", "交易是否成功"], [["2026-01-01", "A", "B", "10", value]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "unknown_status"


@pytest.mark.parametrize("value", ["-d", "de bit", "-收入", "收 入"])
def test_direction_values_cannot_be_repaired_by_deleting_characters(value):
    plan = b(direction=value)
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "invalid_direction"


@pytest.mark.parametrize("value", ["1￥23", "1¥2.3", "￥12￥", "12￥34.56"])
def test_currency_symbols_inside_numeric_token_are_not_extracted(value):
    plan = b(amount=value)
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "invalid_amount"


@pytest.mark.parametrize("value", ["￥1,234.56", "1,234.56¥", "¥ 1234.56", "+￥1234.56"])
def test_whole_currency_prefix_or_suffix_is_accepted(value):
    assert b(amount=value)["rows"][0]["amount"] == Decimal("1234.56")


@pytest.mark.parametrize("unit", ["人民币元", "CNY"])
def test_explicit_unit_currency_cannot_conflict_with_row_currency(unit):
    plan = a(["date", "payer_account", "payee_account", f"amount ({unit})", "currency"], [["2026-01-01", "A", "B", "12", "USD"]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "currency_unit_conflict"


def test_yen_symbol_alone_does_not_establish_cny_currency():
    plan = a(["date", "payer_account", "payee_account", "amount", "currency"], [["2026-01-01", "A", "B", "¥12", "JPY"]])
    assert plan["rows"][0]["currency"] == "JPY"


def test_yen_symbol_is_not_an_unambiguous_currency_column_value():
    plan = a(["date", "payer_account", "payee_account", "amount", "currency"], [["2026-01-01", "A", "B", "12", "¥"]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "invalid_currency"


def _workbook_bytes(workbook):
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_numeric_excel_account_preserves_explicit_zero_padding():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "payer_account", "payee_account", "amount"])
    sheet.append(["2026-01-01", 12345, "B", "12"])
    sheet["B2"].number_format = "0000000000"
    plan = tables.parse(_workbook_bytes(workbook), suffix=".xlsx")
    assert plan["rows"][0]["payer_account"] == "0000012345"


def test_complex_numeric_excel_identity_format_is_not_silently_ignored():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "payer_account", "payee_account", "amount"])
    sheet.append(["2026-01-01", 12345, "B", "12"])
    sheet["B2"].number_format = "000-000"
    plan = tables.parse(_workbook_bytes(workbook), suffix=".xlsx")
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "unsupported_identity_format"


def test_blank_signed_amount_column_does_not_grant_sign_semantics_to_other_column():
    plan = csv_plan(["date", "account", "amount", "signed amount"], [["2026-01-01", "A", "-10", ""]])
    assert not plan["rows"]
    assert "negative_amount" in plan["records"][0]["details"]["validation_issues"]


def test_explicit_midnight_timestamp_is_not_date_only_evidence():
    plan = a(["交易时间", "date", "payer_account", "payee_account", "amount"], [["2026-01-01 00:00:00", "2026-01-01 12:00:00", "A", "B", "10"]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "conflicting_time"


def test_explicit_midnight_cannot_be_overridden_by_separate_clock():
    plan = a(["交易时间", "time of day", "payer_account", "payee_account", "amount"], [["2026-01-01 00:00:00", "12:00:00", "A", "B", "10"]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "conflicting_time"


@pytest.mark.parametrize("own,direction", [("A", "收入"), ("B", "支出"), ("C", "收入")])
def test_bilateral_flow_checks_explicit_own_account_direction(own, direction):
    plan = a(["date", "payer_account", "payee_account", "amount", "本方账号", "direction"], [["2026-01-01", "A", "B", "10", own, direction]])
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "conflicting_direction"


def test_xlsx_dimension_cannot_hide_source_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "payer_account", "payee_account", "amount"])
    sheet.append(["2026-01-01", "A", "B", "10"])
    sheet.append(["2026-01-02", "A", "B", "20"])
    payload = _workbook_bytes(workbook)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as original, zipfile.ZipFile(output, "w") as changed:
        for member in original.infolist():
            data = original.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:D2"', data)
            changed.writestr(member, data)
    plan = tables.parse(output.getvalue(), suffix=".xlsx")
    assert len(plan["rows"]) == 2
    assert len(plan["records"]) == 2


def test_csv_stops_reading_records_when_limit_is_crossed(monkeypatch):
    monkeypatch.setattr(tables, "MAX_RECORDS", 3)
    original = tables.csv.reader

    def bounded(*args, **kwargs):
        for index, row in enumerate(original(*args, **kwargs)):
            assert index <= 3, "parser read past the first over-limit record"
            yield row

    monkeypatch.setattr(tables.csv, "reader", bounded)
    with pytest.raises(tables.TableImportError, match="范围"):
        a(records=[["2026-01-01", "A", "B", "10"] for _ in range(10)])


def test_excel_stops_streaming_at_first_over_limit_record(monkeypatch):
    import openpyxl.worksheet._read_only as read_only
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "payer_account", "payee_account", "amount"])
    for _ in range(10):
        sheet.append(["2026-01-01", "A", "B", "10"])
    payload = _workbook_bytes(workbook)
    monkeypatch.setattr(tables, "MAX_RECORDS", 3)
    original = read_only.ReadOnlyWorksheet.iter_rows

    def bounded(*args, **kwargs):
        for index, row in enumerate(original(*args, **kwargs)):
            assert index <= 3, "parser read past the first over-limit record"
            yield row

    monkeypatch.setattr(read_only.ReadOnlyWorksheet, "iter_rows", bounded)
    with pytest.raises(tables.TableImportError, match="范围"):
        tables.parse(payload, suffix=".xlsx")


@pytest.mark.parametrize("field", ["date", "payer_account", "payee_account", "amount", "currency", "本方账号", "direction"])
def test_formula_in_critical_candidate_is_unresolved_without_execution(field):
    workbook = Workbook()
    sheet = workbook.active
    header = ["date", "payer_account", "payee_account", "amount", "currency", "本方账号", "direction"]
    sheet.append(header)
    values = ["2026-01-01", "A", "B", "10", "CNY", "A", "支出"]
    values[header.index(field)] = '=IF(1=1,"SYN","OTHER")'
    sheet.append(values)
    plan = tables.parse(_workbook_bytes(workbook), suffix=".xlsx")
    assert not plan["rows"]
    assert plan["records"][0]["reason_code"] == "formula_value"


def test_unmapped_excel_formula_does_not_discard_an_otherwise_valid_flow():
    plan = workbook_plan(("Synthetic", [["date", "payer_account", "payee_account", "amount", "无关计算"], ["2026-01-01", "A", "B", "10", "=1+1"]]))
    assert len(plan["rows"]) == 1


def test_workbook_sheet_limit_is_checked_before_loading_source_rows(monkeypatch):
    import openpyxl.worksheet._read_only as read_only
    workbook = Workbook()
    for _ in range(256):
        workbook.create_sheet()
    payload = _workbook_bytes(workbook)

    def forbidden(*args, **kwargs):
        raise AssertionError("must reject sheet count before row iteration")

    monkeypatch.setattr(read_only.ReadOnlyWorksheet, "iter_rows", forbidden)
    with pytest.raises(tables.TableImportError, match="工作表数量"):
        tables.parse(payload, suffix=".xlsx")


@pytest.mark.parametrize("excel", [False, True])
def test_later_reordered_header_cannot_reverse_bilateral_parties(excel):
    first = ["date", "payer_account", "payee_account", "amount"]
    rest = [["2026-01-01", "A", "B", "10"],
            ["date", "payee_account", "payer_account", "amount"],
            ["2026-01-02", "C", "D", "20"]]
    plan = workbook_plan(("Synthetic", [first, *rest])) if excel else csv_plan(first, rest)
    assert [(row["payer_account"], row["payee_account"]) for row in plan["rows"]] == [("A", "B"), ("D", "C")]
    assert [record["reason_code"] for record in plan["records"]] == ["transaction_valid", "repeated_header", "transaction_valid"]
    assert plan["documents"][0]["header_rows"] == [1, 3]
    assert {evidence["header_row"] for evidence in plan["mapping_evidence"]} == {1, 3}
    assert_conserved(plan)


@pytest.mark.parametrize("excel", [False, True])
def test_later_regions_recompute_split_direction_and_monetary_units(excel):
    first = ["date", "account", "income", "expense"]
    rest = [["2026-01-01", "A", "10", "0"],
            ["expense (分)", "date", "account", "income (元)"],
            ["120", "2026-01-02", "A", "0"],
            ["date", "account", "amount (万元)", "direction"],
            ["2026-01-03", "A", "0.000123", "收入"],
            ["2026-01-04", "A", "0.000123", "未知"]]
    plan = workbook_plan(("Synthetic", [first, *rest])) if excel else csv_plan(first, rest)
    assert [row["amount"] for row in plan["rows"]] == [Decimal("10.00"), Decimal("1.20"), Decimal("1.23")]
    assert [row["payer_account"] for row in plan["rows"]] == ["", "A", ""]
    assert plan["records"][-1]["reason_code"] == "invalid_direction"
    assert plan["documents"][0]["header_rows"] == [1, 3, 5]
    assert_conserved(plan)


def test_later_region_without_safe_structure_does_not_reuse_prior_mapping():
    first = ["date", "payer_account", "payee_account", "amount"]
    plan = csv_plan(first, [["2026-01-01", "A", "B", "10"],
                            ["date", "not an account", "unknown", "amount"],
                            ["2026-01-02", "C", "D", "20"]])
    assert len(plan["rows"]) == 1
    assert plan["records"][-1]["reason_code"] == "unknown_table"
    assert_conserved(plan)


def test_explicit_mapping_is_rebound_by_label_in_later_reordered_region():
    first = ["date", "payer_account", "payee_account", "amount"]
    plan = csv_plan(first, [["2026-01-01", "A", "B", "10"],
                            ["date", "payee_account", "payer_account", "amount"],
                            ["2026-01-02", "C", "D", "20"]], mapping={"payer_account": "payer_account", "payee_account": "payee_account"})
    assert [(row["payer_account"], row["payee_account"]) for row in plan["rows"]] == [("A", "B"), ("D", "C")]


def test_explicit_mapping_missing_in_later_region_is_not_forced_by_position():
    first = ["date", "payer_account", "payee_account", "amount"]
    plan = csv_plan(first, [["2026-01-01", "A", "B", "10"],
                            ["date", "to account", "from account", "amount"],
                            ["2026-01-02", "C", "D", "20"]], mapping={"payer_account": "payer_account", "payee_account": "payee_account"})
    assert len(plan["rows"]) == 1
    assert plan["records"][-1]["reason_code"] == "region_mapping_conflict"
    assert_conserved(plan)


def test_later_region_ambiguous_money_is_validated_without_first_region_column_count():
    first = ["date", "payer_account", "payee_account", "amount"]
    plan = csv_plan(first, [["2026-01-01", "A", "B", "10"],
                            ["date", "payer_account", "payee_account", "amount", "发生金额"],
                            ["2026-01-02", "C", "D", "20", "30"]])
    assert len(plan["rows"]) == 1
    assert plan["records"][-1]["reason_code"] == "conflicting_amount"
    assert_conserved(plan)


@pytest.mark.parametrize("amount", ["12", ""])
def test_missing_time_data_with_header_words_cannot_replace_current_region(amount):
    first = ["date", "payer_account", "payee_account", "amount", "summary"]
    plan = csv_plan(first, [["2026-01-01", "A", "B", "10", "normal"],
                            ["", "A", "date", amount, "amount"],
                            ["2026-01-02", "C", "D", "20", "normal"]])
    assert [(row["payer_account"], row["payee_account"]) for row in plan["rows"]] == [("A", "B"), ("C", "D")]
    assert plan["records"][1]["disposition"] == "unresolved"
    assert plan["records"][1]["reason_code"] == "missing_time"
    assert plan["documents"][0]["header_rows"] == [1]
    assert_conserved(plan)
