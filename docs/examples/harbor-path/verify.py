"""Check the public synthetic input against files downloaded through the UI."""
import csv
from decimal import Decimal
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def verify():
    with (ROOT / 'fundstrace-synthetic.csv').open(encoding='utf-8-sig', newline='') as handle:
        source = list(csv.DictReader(handle))
    workbook = load_workbook(ROOT / 'fundstrace-synthetic.xlsx', data_only=True)
    rows = list(workbook.active.values)
    assert source == [dict(zip(rows[0], row)) for row in rows[1:]], 'CSV and XLSX differ'
    workbook.close()
    exported = load_workbook(ROOT / 'exports/tracked-flows.xlsx', data_only=True)
    assert exported.sheetnames == ['追踪流水', '待调取账号']
    actual_rows = list(exported['追踪流水'].values)
    assert len(actual_rows) - 1 == 12, 'Unexpected extra or duplicate export rows'
    actual = {row[1]: dict(zip(actual_rows[0], row)) for row in actual_rows[1:]}
    assert len(source) == len(actual) == 12
    fields = {'本方户名': '付款户名', '本方账号': '付款账号',
              '对方户名': '收款户名', '对方账号': '收款账号',
              '交易金额': '交易金额', '币种': '币种', '摘要': '摘要'}
    for expected in source:
        row = actual[expected['交易流水号']]
        assert expected['收支标志'] == '支出'
        assert row['交易时间'].strftime('%Y-%m-%d %H:%M:%S') == expected['交易时间']
        for before, after in fields.items():
            assert row[after] == expected[before], (expected['交易流水号'], before)
        assert isinstance(row['交易金额'], str), 'Export amount must retain its decimal string'
    exported.close()
    accounts = {r[key] for r in source for key in ('本方账号', '对方账号')}
    pairs = {(r['本方账号'], r['对方账号']) for r in source}
    total = sum((Decimal(r['交易金额']) for r in source), Decimal('0.00'))
    assert len(accounts) == 7 and len(pairs) == 8 and total == Decimal('1832503.00')
    anchor = source[0]
    extended = [r for r in source if r['本方账号'] == anchor['对方账号']
                and r['交易时间'] > anchor['交易时间']][:5]
    assert len(extended) == 4
    assets = ['fundstrace-synthetic.csv', 'fundstrace-synthetic.xlsx',
              'exports/tracked-flows.xlsx', 'exports/traced-graph.png']
    return {'syntheticOnly': True, 'fileChecks': 'passed', 'transactions': len(source),
            'accounts': len(accounts), 'directedRelations': len(pairs),
            'transactionAmountSumByCurrency': {'CNY': format(total, '.2f')},
            'expectedFirstAnchorTrackedCount': 1 + len(extended),
            'sha256': {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in assets}}


if __name__ == '__main__':
    print(json.dumps(verify(), ensure_ascii=False, indent=2))
