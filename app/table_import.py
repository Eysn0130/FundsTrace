"""Versioned, deterministic CSV/workbook interpretation, without persistent writes.

Column names identify candidates, never a confidence score. Every logical source
record has a disposition; unsafe monetary interpretations remain source records.
The uploaded file is the original evidence, rather than a second raw-cell table.
"""
from __future__ import annotations

import csv
import datetime as dt
from dataclasses import dataclass
from decimal import Decimal, DecimalException, localcontext
import hashlib
import io
import json
from pathlib import Path
import re
import unicodedata
import zipfile

from .importer import FIELD_SYNONYMS, _row_key
from .table_mapping import identity_field
from . import table_financial, table_contracts, table_time_formats, table_source_roles


IDENTITY_PROFILE = "funds-trace-table-v5"
ROLE_PROFILE = "funds-trace-table-v6"
FINANCIAL_PROFILE = "funds-trace-table-v7"
CONTRACT_PROFILE = "funds-trace-table-v8"
DIRECTION_PROFILE = "funds-trace-table-v9"
TIME_FORMAT_PROFILE = "funds-trace-table-v10"
UNIT_DECLARATION_PROFILE = "funds-trace-table-v11"
SOURCE_ROLE_PROFILE = "funds-trace-table-v12"
BANK_NAME_PROFILE = "funds-trace-table-v13"
QUERY_OWNER_PROFILE = "funds-trace-table-v14"
REGION_HEADER_PROFILE = "funds-trace-table-v15"
REORDERED_TIME_PROFILE = "funds-trace-table-v16"
PROFILE = REORDERED_TIME_PROFILE
REGION_HEADER_PROFILES = frozenset((REGION_HEADER_PROFILE, REORDERED_TIME_PROFILE))
QUERY_OWNER_PROFILES = frozenset((QUERY_OWNER_PROFILE, *REGION_HEADER_PROFILES))
BANK_NAME_PROFILES = frozenset((BANK_NAME_PROFILE, *QUERY_OWNER_PROFILES))
SOURCE_ROLE_PROFILES = frozenset((SOURCE_ROLE_PROFILE, *BANK_NAME_PROFILES))
UNIT_DECLARATION_PROFILES = frozenset((UNIT_DECLARATION_PROFILE, *SOURCE_ROLE_PROFILES))
TIME_FORMAT_PROFILES = frozenset((TIME_FORMAT_PROFILE, *UNIT_DECLARATION_PROFILES))
DIRECTION_PROFILES = frozenset((DIRECTION_PROFILE, *TIME_FORMAT_PROFILES))
CONTRACT_PROFILES = frozenset((CONTRACT_PROFILE, *DIRECTION_PROFILES))
FINANCIAL_PROFILES = frozenset((FINANCIAL_PROFILE, *CONTRACT_PROFILES))
ROLE_PROFILES = frozenset((ROLE_PROFILE, *FINANCIAL_PROFILES))
SUPPORTED_PROFILES = frozenset((IDENTITY_PROFILE, *ROLE_PROFILES))
ROLE_CONFLICT_BASIS = "显式身份角色或候选字段存在矛盾"
HISTORY_CONFLICT_BASIS = "同一物理来源已有不同资金解释，保留历史事实，防止重复累计"
FINANCIAL_CONFLICT_BASIS = "显式资金字段或单位声明未通过全量校验"
FINANCIAL_VALIDATION_ISSUES = frozenset((
    "missing_time", "invalid_time", "conflicting_time", "missing_amount", "invalid_amount",
    "amount_precision", "amount_overflow", "conflicting_amount", "unknown_unit", "conflicting_unit",
    "invalid_direction", "conflicting_direction", "conflicting_income_expense", "invalid_balance",
    "currency_unit_conflict", "invalid_currency", "unknown_status", "formula_value",
    "balance_amount_conflict", "balance_direction_conflict",
))
ROW_KIND = "table_record"
TRANSACTION_REFERENCE_POLICY = "transaction-reference-v1"
BANK_FEEDBACK_32_COLUMNS = frozenset("交易卡号 交易账号 查询反馈结果原因 交易户名 交易证件号码 交易时间 交易金额 交易余额 收付标志 交易对手账卡号 现金标志 对手户名 对手身份证号 对手开户银行 摘要说明 交易币种 交易网点名称 交易发生地 交易是否成功 传票号 IP地址 MAC地址 对手交易余额 交易流水号 日志号 凭证种类 凭证号 交易柜员号 备注 商户名称 商户代码 交易类型".split())
MAX_BYTES = 256 * 1024 * 1024
MAX_RECORDS = 1_000_000
MAX_COLUMNS = 256
MAX_FIELD = 1024 * 1024
MAX_SHEETS = 256
_MAX_MONEY = Decimal("9999999999999999.99")
_CENT = Decimal("0.01")
REASONS = {
    "transaction_valid": "资金字段已通过全量校验",
    "unknown_table": "表格结构尚不能唯一识别，未计入资金统计",
    "missing_time": "缺少交易时间，未计入资金统计",
    "invalid_time": "交易时间格式无效或日期含义不明确，未计入资金统计",
    "conflicting_time": "多个时间字段不一致，未计入资金统计",
    "missing_amount": "缺少交易金额，未计入资金统计",
    "invalid_amount": "金额格式无法精确解析，未计入资金统计",
    "amount_precision": "换算后的金额不能按分精确表示，未计入资金统计",
    "amount_overflow": "金额超出存储范围，未计入资金统计",
    "conflicting_amount": "多个金额字段换算后不一致，未计入资金统计",
    "unknown_unit": "金额单位不能唯一确定，未计入资金统计",
    "conflicting_unit": "金额单位声明互相矛盾，未计入资金统计",
    "negative_amount": "负号含义没有明确依据，未计入资金统计",
    "missing_direction": "缺少可确定的收支方向，未计入资金统计",
    "invalid_direction": "方向枚举尚未适配，未计入资金统计",
    "conflicting_direction": "方向与金额或收支分列矛盾，未计入资金统计",
    "conflicting_income_expense": "收入与支出同时非零，未计入资金统计",
    "missing_identity": "缺少必要的本方或双边身份，未计入资金统计",
    "conflicting_field": "同一字段的多个候选列不一致，未计入资金统计",
    "invalid_currency": "币种不能唯一识别，未计入资金统计",
    "invalid_balance": "余额不能精确解释，未计入资金统计",
    "balance_amount_conflict": "本笔交易前后账面余额差与金额不一致，未计入资金统计",
    "balance_direction_conflict": "本笔交易前后账面余额变化与收支方向矛盾，未计入资金统计",
    "row_width": "来源记录列数与表头不同，未计入资金统计",
    "zero_amount": "零金额记录不计入资金统计",
    "blank_record": "空白记录不计入资金统计",
    "preamble": "表头前说明记录不计入资金统计",
    "repeated_header": "重复表头不计入资金统计",
    "unknown_status": "交易状态尚未适配，未计入资金统计",
    "failed_transaction": "来源明确标记交易未成功，不计入资金统计",
    "feedback_only": "查询反馈说明不属于交易记录",
    "unsafe_account_number": "Excel数值账号可能已损失精度，未计入资金统计",
    "auxiliary_table": "辅助表仅保留来源处置元数据，不计入资金统计",
    "column_count_mismatch": "来源记录列数与表头不同，未计入资金统计",
    "currency_unit_conflict": "表头明示的金额币种与来源币种矛盾，未计入资金统计",
    "unsupported_identity_format": "Excel身份数字显示格式无法唯一解释，未计入资金统计",
    "formula_value": "关键字段为未执行的Excel公式，未计入资金统计",
    "region_mapping_conflict": "所选映射不能唯一适用于当前表格区域，未计入资金统计",
}


class TableImportError(ValueError):
    """Only safe, constant diagnostics; never interpolate source cells."""


def _norm(value):
    value = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return re.sub(r"[\s_\-]+", "", value)


def _enum(value):
    # Header normalization is deliberately more permissive than value parsing.
    # Never delete signs or internal characters from a financial enum value.
    return unicodedata.normalize("NFKC", _text(value)).strip().lower()


def _unit_token(value):
    # Unit values are financial enums, not header names. Keep all signs and
    # punctuation; whitespace alone may separate the known English words.
    return re.sub(r"\s+", "", _enum(value))


_EXTRA = {
    "trade_no": ["transaction id", "transaction number", "reference", "reference number", "trade_no"],
    "trade_time": ["transaction date", "transaction time", "transaction datetime", "datetime", "timestamp", "posting date", "trade_time"],
    "time_part": ["交易时分秒", "交易时刻", "时分秒", "time of day"],
    "payer_name": ["payer name", "sender name", "from name", "payer_name", "付款方户名"],
    "payer_account": ["payer account", "sender account", "from account", "debit account", "payer_account"],
    "payer_bank": ["payer bank", "sender bank", "payer_bank"],
    "payee_name": ["payee name", "recipient name", "to name", "payee_name", "收款方户名"],
    "payee_account": ["payee account", "recipient account", "to account", "credit account", "payee_account"],
    "payee_bank": ["payee bank", "recipient bank", "payee_bank"],
    "self_name": ["account name", "own name", "self_name", "交易户名"],
    "self_account": ["account", "account number", "own account", "self account", "self_account", "本方帐号", "帐号"],
    "counter_name": ["counterparty name", "counterparty", "counter_name"],
    "counter_account": ["counterparty account", "counter account", "counter_account", "对方帐号", "对手帐号"],
    "self_card": ["交易卡号"],
    "counter_bank": ["counterparty bank", "counter_bank", "对手开户银行"],
    "direction": ["credit/debit", "debit/credit", "dr/cr", "credit debit indicator", "收支方向"],
    "amount": ["transaction amount", "transfer amount", "signed amount", "带符号金额", "带符号交易金额"],
    "income_amt": ["income", "credit", "credit amount", "money in", "deposit", "收入", "贷方发生金额"],
    "expense_amt": ["expense", "debit", "debit amount", "money out", "withdrawal", "支出", "借方发生金额"],
    "balance": ["closing balance", "running balance", "交易余额"],
    "currency": ["currency code", "ccy", "币别"],
    "summary": ["description", "memo", "narrative", "摘要说明"],
    "channel": ["channel", "transaction type"],
    "status": ["交易是否成功", "交易状态", "transaction status", "status"],
    "feedback_reason": ["查询反馈结果原因"],
}
_ALIASES = {}
for _field in set(FIELD_SYNONYMS) | set(_EXTRA):
    _ALIASES[_field] = {_norm(x) for x in [*FIELD_SYNONYMS.get(_field, []), *_EXTRA.get(_field, [])]}
# Cash/remittance indicators are not currency codes.
_ALIASES["currency"].discard(_norm("钞汇标志"))
_FINANCIAL_ALIASES = {field: {_norm(value) for value in values}
                      for field, values in table_financial.ALIASES.items()}
_FINANCIAL_MONEY_HEADERS = frozenset(value for field in table_financial.MONEY_FIELDS
                                    for value in _ALIASES[field] | _FINANCIAL_ALIASES.get(field, set()))
_UNIT = {
    "元": Decimal(1), "人民币元": Decimal(1), "yuan": Decimal(1), "cny": Decimal(1),
    "分": Decimal("0.01"), "cent": Decimal("0.01"), "cents": Decimal("0.01"),
    "角": Decimal("0.1"), "jiao": Decimal("0.1"),
    "千元": Decimal(1000), "thousandyuan": Decimal(1000),
    "万元": Decimal(10000), "tenthousandyuan": Decimal(10000),
}
_SIGNED = {_norm(x) for x in ("signed amount", "带符号金额", "带符号交易金额")}
_IN = frozenset(("c", "cr", "crdt", "credit", "in", "income", "收", "入", "进", "贷", "收入", "转入", "收款", "贷方", "贷记", "收方", "存入", "入账", "进账", "存现"))
_OUT = frozenset(("d", "dr", "dbit", "debit", "out", "expense", "支", "出", "付", "借", "支出", "转出", "付款", "借方", "借记", "付方", "取现", "取款", "出账", "支取", "提现"))
_CURRENCY = {"人民币": "CNY", "人民币元": "CNY", "RMB": "CNY", "美元": "USD", "美金": "USD", "欧元": "EUR", "港币": "HKD", "港元": "HKD", "日元": "JPY", "英镑": "GBP", "澳元": "AUD", "加元": "CAD", "新加坡元": "SGD", "新西兰元": "NZD", "瑞士法郎": "CHF"}
_ISO_CURRENCIES = frozenset({"CNY", "USD", "EUR", "HKD", "JPY", "GBP", "AUD", "CAD", "SGD", "NZD", "CHF", "KRW", "THB", "MYR", "MOP", "TWD", "RUB", "INR", "BRL", "ZAR", "AED", "SAR", "SEK", "NOK", "DKK", "PLN", "CZK", "HUF", "TRY", "IDR", "PHP", "VND", "MXN", "ILS"})
_CONTRACT_ALIASES = {field: {_norm(value) for value in values}
                     for field, values in table_contracts.ALIASES.items()}
_DIRECTION_HEADERS = frozenset(_norm(value) for value in (
    "debit/credit indicator", "credit/debit indicator", "cr/dr",
    "借贷标识", "收支标识", "收付标识"))
_CLOCK_HEADERS = frozenset(_norm(value) for value in (
    "交易时间", "业务时间", "记账时间", "转账时间", "发生时间", "时间",
    "time", "transaction time", "trade time"))
_AFTER_HEADERS = {_norm(value) for value in table_contracts.AFTER_HEADERS}
_BEFORE_HEADERS = {_norm(value) for value in table_contracts.BEFORE_HEADERS}
_CONTRACT_MONEY_HEADERS = frozenset(value for field in (*table_financial.MONEY_FIELDS, *table_contracts.MONEY_FIELDS)
                                    for value in _CONTRACT_ALIASES.get(field, ()))
FINANCIAL_MAPPING_FIELDS = frozenset((*_ALIASES, *_FINANCIAL_ALIASES))
MAPPING_FIELDS = FINANCIAL_MAPPING_FIELDS | frozenset(_CONTRACT_ALIASES) | {"counter_card"}
REASON_CODES = frozenset(REASONS)
DETAILS_FIELDS = frozenset({"sheet", "source_row", "financial_effect", "reason", "validation_issues", "validation_basis", "currency_basis", "counterparty_basis", "sign_convention", "flow_sha256", "source_account_column", "reported_status", "source_role"})
REFERENCE_DETAILS_FIELDS = frozenset({"trade_no_basis"})
FLOW_FIELDS = ("trade_no", "trade_time", "payer_name", "payer_account", "payer_bank", "payee_name", "payee_account", "payee_bank", "amount", "balance", "currency", "summary", "channel")


def mapping_fields(profile):
    """A new identity role must not become an ignored old-version override."""
    return MAPPING_FIELDS if profile in SOURCE_ROLE_PROFILES else MAPPING_FIELDS - {"counter_card"}


def record_key(path, member_sha, source_row):
    """Stable identity for one physical table occurrence under this parser."""
    return hashlib.sha256(json.dumps([IDENTITY_PROFILE, path, member_sha, source_row], separators=(",", ":")).encode()).hexdigest()


def flow_digest(row):
    """Hash the established deduplication facts; never include mutable source IDs."""
    values = list(_row_key(row))
    values[8] = format(values[8], ".2f")
    values[9] = None if values[9] is None else format(values[9], ".2f")
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


@dataclass(frozen=True, slots=True)
class _WorkbookValue:
    value: object
    number_format: str = "General"
    formula: bool = False
    date_only: bool = False


def _raw(value):
    return value.value if isinstance(value, _WorkbookValue) else value


def _text(value):
    value = _raw(value)
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat(sep=" ") if isinstance(value, dt.datetime) else value.isoformat()
    return str(value).strip()


def _header_parts(value, profile=IDENTITY_PROFILE):
    value = unicodedata.normalize("NFKC", _text(value))
    if profile in FINANCIAL_PROFILES:
        normalized = _norm(value)
        money_headers = _FINANCIAL_MONEY_HEADERS | (_CONTRACT_MONEY_HEADERS if profile in CONTRACT_PROFILES else frozenset())
        if normalized in _ALIASES["direction"] or profile in DIRECTION_PROFILES and normalized in _DIRECTION_HEADERS:
            return normalized, Decimal(1), "元/默认基础单位"
        if profile in UNIT_DECLARATION_PROFILES:
            # A complete monetary field may declare an existing unit in either
            # language. Pair delimiters before interpreting the declaration;
            # the historical fallback otherwise accepts mixed brackets.
            paired = re.fullmatch(r"(.*?)\s*(?:\(([^()\[\]]+)\)|\[([^()\[\]]+)\])\s*", value)
            if paired and _norm(paired[1]) in money_headers:
                declared = re.sub(r"^(?:单位\s*[:：]?\s*|(?i:unit)\s*[:：]\s*)", "", (paired[2] or paired[3]).strip())
                token = _unit_token(declared)
                return _norm(paired[1]), _UNIT.get(token), token
            malformed = re.fullmatch(r"(.*?)\s*([\(\[].*)", value)
            if malformed and _norm(malformed[1]) in money_headers:
                return _norm(malformed[1]), None, _unit_token(malformed[2])
        # Only complete monetary names may declare a unit. A slash in a
        # direction such as debit/credit must keep its original meaning.
        declaration = re.fullmatch(r"(.*?)\s*[\(\[]\s*(?:单位\s*[:：]?\s*)?([^()\[\]]+)\s*[\)\]]\s*", value)
        if not declaration:
            declaration = re.fullmatch(r"(.*?)\s*/\s*(?:单位\s*[:：]?\s*)?([^/]+)\s*", value)
        if declaration and _norm(declaration[1]) in money_headers:
            declared = declaration[2]
            if profile in UNIT_DECLARATION_PROFILES:
                declared = re.sub(r"^(?i:unit)\s*[:：]\s*", "", declared.strip())
            token = _unit_token(declared)
            return _norm(declaration[1]), _UNIT.get(token), token
        if profile in UNIT_DECLARATION_PROFILES:
            malformed_slash = re.fullmatch(r"([^/]+)/(.*)", value)
            if malformed_slash and _norm(malformed_slash[1]) in money_headers:
                return _norm(malformed_slash[1]), None, _unit_token(malformed_slash[2])
        normalized = _unit_token(value)
        for unit in sorted(_UNIT, key=len, reverse=True):
            base = _norm(normalized[:-len(unit)])
            if normalized.endswith(unit) and base in money_headers:
                return base, _UNIT[unit], unit
        # A recognizable monetary declaration with an invalid unit must not
        # fall back to historical header normalization that deletes its signs.
        legacy = _header_parts(value, IDENTITY_PROFILE)
        if legacy[2] != "元/默认基础单位" and legacy[0] in money_headers:
            return legacy[0], None, normalized
        return legacy
    match = re.fullmatch(r"(.*?)\s*[（(\[]\s*(?:单位\s*[:：]?\s*)?([^()（）\[\]]+)\s*[）)\]]\s*", value)
    if match:
        base, unit = _norm(match[1]), _norm(match[2])
        # Parentheses which are not monetary units must not erase semantics.
        if unit in _UNIT or base in _ALIASES["amount"] | _ALIASES["balance"] | _ALIASES["income_amt"] | _ALIASES["expense_amt"]:
            return base, _UNIT.get(unit), unit
    # Explicit unit suffix, such as amount_cents and 交易金额万元.
    normalized = _norm(value)
    for unit in sorted(_UNIT, key=len, reverse=True):
        if normalized.endswith(unit) and normalized[:-len(unit)] in (_ALIASES["amount"] | _ALIASES["balance"] | _ALIASES["income_amt"] | _ALIASES["expense_amt"]):
            return normalized[:-len(unit)], _UNIT[unit], unit
    return normalized, Decimal(1), "元/默认基础单位"


def _declared_time_header(value, profile):
    if profile not in TIME_FORMAT_PROFILES:
        return None
    declaration = table_time_formats.header_parts(_text(value))
    if declaration is None:
        return None
    base, spec = declaration
    base = _norm(base)
    clocks = _ALIASES["time_part"] | _FINANCIAL_ALIASES.get("time_part", set())
    dates = _ALIASES["trade_time"] | _FINANCIAL_ALIASES.get("trade_time", set())
    if spec in {"HH:MM:SS", "HHMMSS"}:
        return (base, spec) if base in clocks | _CLOCK_HEADERS else None
    return (base, spec) if base in dates else None


def _columns(headers):
    counts = {}
    labels = []
    for index, raw in enumerate(headers):
        label = _text(raw) or f"未命名列{index + 1}"
        counts[label] = counts.get(label, 0) + 1
        labels.append(label)
    return [f"{label} [列{i + 1}]" if counts[label] > 1 else label for i, label in enumerate(labels)]


def _candidates(headers, profile=IDENTITY_PROFILE):
    found = {}
    query_owner = (profile in QUERY_OWNER_PROFILES and table_source_roles.classify(
        [_text(header) for header in headers], query_conditions=True)
        in {"transaction_query_only", "transaction_explicit"})
    for index, header in enumerate(headers):
        if query_owner and table_source_roles.query_condition(_text(header)):
            continue
        base, factor, unit = _header_parts(header, profile)
        declared_time = _declared_time_header(header, profile)
        if declared_time:
            base = declared_time[0]
        fields = [field for field, synonyms in _ALIASES.items() if base in synonyms]
        basis = "exact_header_alias"
        if profile in FINANCIAL_PROFILES:
            extra = [field for field, synonyms in _FINANCIAL_ALIASES.items() if base in synonyms]
            if not fields and extra:
                fields, basis = extra, "explicit_financial_header"
            elif fields and _header_parts(header) != (base, factor, unit):
                basis = "explicit_financial_header"
        if profile in CONTRACT_PROFILES:
            for field, aliases in _CONTRACT_ALIASES.items():
                if base in aliases and field not in fields:
                    fields.append(field)
                    basis = "explicit_source_contract"
        if profile in DIRECTION_PROFILES and _norm(_text(header)) in _DIRECTION_HEADERS:
            fields, basis = ["direction"], "explicit_financial_header"
        if profile in SOURCE_ROLE_PROFILES:
            # An account and its card are different identifiers, not competing
            # values of one field. Generic query conditions confer no role.
            identity_roles = {"本方卡号": "self_card", "交易卡号": "self_card",
                              "交易对方卡号": "counter_card", "交易对方账卡号": "counter_account"}
            if base in identity_roles:
                fields, basis = [identity_roles[base]], "explicit_identity_role"
        if query_owner and _norm(header) in {"查询帐号", "查询卡号"}:
            fields = ["self_account" if _norm(header) == "查询帐号" else "self_card"]
            basis = "uniform_query_owner"
        if profile in BANK_NAME_PROFILES and _norm(header) == "交易对方账号开户行":
            fields, basis = ["counter_bank"], "explicit_identity_role"
        if declared_time:
            basis = "explicit_time_format"
        if not fields and profile in ROLE_PROFILES and (field := identity_field(base)):
            fields = [field]
            basis = "role_qualified_header"
        for field in fields:
            found.setdefault(field, []).append({"index": index, "factor": factor, "unit": unit, "signed": base in _SIGNED, "unit_currency": "CNY" if unit in {"人民币元", "cny"} else "", "basis": basis, **({"same_entry_after": field == "balance" and base in _AFTER_HEADERS,
                                  "same_entry_before": field == "balance_before" and base in _BEFORE_HEADERS}
                                 if profile in CONTRACT_PROFILES else {}),
                                 **({"time_format": declared_time[1]} if declared_time else {}),
                                 **({"identity_kind": "card" if field == "self_card" else "account",
                                     "contract_owner": True}
                                    if query_owner and field in {"self_account", "self_card"} else {})})
    return found


def _shape(candidates, profile=IDENTITY_PROFILE):
    money = any(candidates.get(k) for k in ("amount", "income_amt", "expense_amt"))
    if not candidates.get("trade_time") or not money:
        return None
    if (candidates.get("payer_account") or candidates.get("payer_name")) and (candidates.get("payee_account") or candidates.get("payee_name")):
        return "A"
    if candidates.get("self_account") or profile in SOURCE_ROLE_PROFILES and candidates.get("self_card"):
        return "B"
    return None


def _locate_header(raw, profile=IDENTITY_PROFILE):
    for index, row in enumerate(raw):
        candidates = _candidates(row, profile)
        if candidates.get("trade_time") and any(candidates.get(k) for k in ("amount", "income_amt", "expense_amt")):
            return index, candidates, _shape(candidates, profile)
    return None, {}, None


def _bytes(source):
    try:
        if isinstance(source, (bytes, bytearray)):
            raw = bytes(source)
        elif hasattr(source, "read"):
            source.seek(0)
            raw = source.read(MAX_BYTES + 1)
        else:
            with Path(source).open("rb") as stream:
                raw = stream.read(MAX_BYTES + 1)
    except (OSError, ValueError, TypeError):
        raise TableImportError("无法读取表格文件") from None
    if not isinstance(raw, bytes) or len(raw) > MAX_BYTES:
        raise TableImportError("表格文件超出安全处理大小")
    return raw


def _csv(raw, suffix, profile=IDENTITY_PROFILE):
    if b"\x00" in raw[:4096] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise TableImportError("CSV编码无法安全识别")
    encodings = ("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
    text, encoding = None, None
    for candidate in encodings:
        try:
            text, encoding = raw.decode(candidate, errors="strict"), candidate
            break
        except UnicodeError:
            continue
    if text is None or "\x00" in text:
        raise TableImportError("CSV编码无法安全识别")
    csv.field_size_limit(max(csv.field_size_limit(), MAX_FIELD))
    choices = []
    for delimiter in ("\t", ",", ";", "|"):
        sample = []
        try:
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            for index, row in enumerate(reader):
                if index >= MAX_RECORDS:
                    raise TableImportError("表格记录数或列数超出安全处理范围")
                sample.append(row)
                if index >= 99:
                    break
        except csv.Error:
            continue
        header, candidates, mode = _locate_header(sample, profile)
        if header is not None:
            choices.append((delimiter, header, candidates, mode))
    if len(choices) > 1:
        # No order-based tie breaking when two separators yield valid schemas.
        raise TableImportError("CSV分隔符存在歧义，无法安全识别")
    if choices:
        delimiter = choices[0][0]
    else:
        # An unknown schema still gets a complete disposition, not a guessed map.
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = "\t" if suffix == ".tsv" else ","
    try:
        rows = _collect_rows(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True))
    except csv.Error:
        raise TableImportError("CSV引号或记录结构损坏，文件未接纳") from None
    return [{"path": "table.csv", "name": "", "rows": rows, "sha256": hashlib.sha256(raw).hexdigest(), "size": len(raw), "encoding": encoding, "delimiter": delimiter}]


def _check_row(row, count):
    if count > MAX_RECORDS or len(row) > MAX_COLUMNS:
        raise TableImportError("表格记录数或列数超出安全处理范围")
    if any(len(_text(value)) > MAX_FIELD for value in row):
        raise TableImportError("表格字段超出安全处理大小")


def _collect_rows(iterator, *, previous_count=0):
    rows = []
    for row in iterator:
        row = list(row)
        _check_row(row, previous_count + len(rows) + 1)
        rows.append(row)
    return rows


def _excel(raw, suffix):
    if suffix == ".xls":
        raise TableImportError("当前环境未提供旧版XLS读取组件，请使用XLSX或CSV格式")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            info = archive.infolist()
            if len(info) > 2048 or sum(member.file_size for member in info) > 300 * 1024 * 1024:
                raise TableImportError("工作簿展开大小超出安全处理范围")
        from openpyxl import load_workbook
        from openpyxl.styles.numbers import is_datetime
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        sheets = []
        try:
            digest = hashlib.sha256(raw).hexdigest()
            count = 0
            if len(workbook.worksheets) > MAX_SHEETS:
                raise TableImportError("工作表数量超出安全处理范围")
            for index, sheet in enumerate(workbook.worksheets, start=1):
                # Read the actual XML stream; producer dimensions may understate
                # rows/columns and must not authorize silent source truncation.
                sheet.reset_dimensions()

                def values():
                    for cells in sheet.iter_rows(values_only=False):
                        row = []
                        for cell in cells:
                            value = cell.value
                            formula = getattr(cell, "data_type", "") == "f"
                            if formula or isinstance(value, (int, float, dt.date, dt.time)):
                                fmt = getattr(cell, "number_format", "General")
                                date_only = (isinstance(value, dt.date) and not isinstance(value, dt.datetime)) or (isinstance(value, dt.datetime) and value.time() == dt.time() and is_datetime(fmt) == "date")
                                value = _WorkbookValue(value, fmt, formula, date_only)
                            row.append(value)
                        yield row

                rows = _collect_rows(values(), previous_count=count)
                width = max((len(row) for row in rows), default=0)
                # XLSX omits trailing empty cells; unlike malformed CSV widths,
                # missing XML cells are a defined empty worksheet coordinate.
                for row in rows:
                    row.extend([None] * (width - len(row)))
                count += len(rows)
                path = f"worksheets/{index:04d}"
                member_sha = hashlib.sha256(json.dumps([digest, index, sheet.title], ensure_ascii=False).encode()).hexdigest()
                sheets.append({"path": path, "name": sheet.title, "rows": rows, "sha256": member_sha, "size": 0, "encoding": "xlsx", "delimiter": ""})
        finally:
            workbook.close()
        return sheets
    except TableImportError:
        raise
    except Exception:
        raise TableImportError("Excel工作簿损坏、加密或格式不受支持") from None


def _money(value, candidate):
    raw = _text(value)
    if not raw or raw.lower() in {"null", "none", "nan", "--", "-"}:
        return None, None
    if candidate["factor"] is None:
        return None, "unknown_unit"
    parentheses = raw.startswith("(") and raw.endswith(")")
    if parentheses:
        raw = raw[1:-1].strip()
        if raw.startswith(("+", "-")):
            return None, "invalid_amount"
    # Currency decoration may surround the entire number; embedded symbols are
    # not characters which can safely be deleted from a numeric token.
    if raw.startswith(("￥", "¥")):
        raw = raw[1:].strip()
    elif raw.endswith(("￥", "¥")):
        raw = raw[:-1].strip()
    elif len(raw) > 1 and raw[0] in "+-" and raw[1] in "￥¥":
        raw = raw[0] + raw[2:].strip()
    if "," in raw or "，" in raw:
        raw = raw.replace("，", ",")
        if not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", raw):
            return None, "invalid_amount"
        raw = raw.replace(",", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", raw):
        return None, "invalid_amount"
    if len(raw) > 256:
        return None, "invalid_amount"
    exponent = re.search(r"[eE]([+-]?\d+)\Z", raw)
    if exponent and abs(int(exponent[1])) > 1000:
        # Bound the exponent before arithmetic: Decimal context underflow must
        # never turn a positive source value into a zero-movement decision.
        return None, "amount_precision" if int(exponent[1]) < 0 else "amount_overflow"
    try:
        with localcontext() as context:
            context.prec = max(50, len(raw) + 12)
            amount = Decimal(raw) * candidate["factor"]
            if parentheses:
                amount = -amount
            if not amount.is_finite() or abs(amount) > _MAX_MONEY:
                return None, "amount_overflow"
            if amount != amount.quantize(_CENT):
                return None, "amount_precision"
            return amount.quantize(_CENT), None
    except (DecimalException, OverflowError, ValueError):
        return None, "invalid_amount"


def _time(value):
    value = _raw(value)
    if isinstance(value, dt.datetime):
        return value if value.tzinfo is None else None
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time())
    raw = _text(value)
    raw = raw.replace("年", "-").replace("月", "-").replace("日", "")
    # Year-first dates have an unambiguous day/month convention. No inference.
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T]?(\d{1,2}):(\d{2})(?::(\d{2})(?:\.(\d{1,6}))?)?)?", raw)
    if match:
        parts = match.groups()
        try:
            return dt.datetime(int(parts[0]), int(parts[1]), int(parts[2]), int(parts[3] or 0), int(parts[4] or 0), int(parts[5] or 0), int((parts[6] or "").ljust(6, "0") or 0))
        except ValueError:
            return None
    for fmt in ("%Y%m%d%H%M%S", "%Y%m%d %H%M%S", "%Y%m%d"):
        if re.fullmatch(r"\d{8}|\d{14}|\d{8} \d{6}", raw):
            try:
                return dt.datetime.strptime(raw, fmt)
            except ValueError:
                continue
    return None


def _date_only(value):
    if isinstance(value, _WorkbookValue) and value.date_only:
        return True
    value = _raw(value)
    if isinstance(value, dt.datetime):
        return False
    if isinstance(value, dt.date):
        return True
    raw = _text(value).replace("年", "-").replace("月", "-").replace("日", "")
    return bool(re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{8}", raw))


def _value(row, candidate):
    return row[candidate["index"]] if candidate["index"] < len(row) else ""


def _field(row, candidates, name, issues, *, normalize=None):
    values = []
    for candidate in candidates.get(name, []):
        value = _value(row, candidate)
        if name in {"payer_account", "payee_account", "self_account", "self_card", "counter_account", "counter_card"} and isinstance(value, _WorkbookValue):
            raw = value.value
            if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                try:
                    integer = int(raw)
                    if raw != integer or integer < 0 or len(str(integer)) > 15:
                        issues.append("unsafe_account_number")
                    elif value.number_format not in {"General", "0", "@"}:
                        if re.fullmatch(r"0+", value.number_format):
                            value = str(integer).zfill(len(value.number_format))
                        else:
                            issues.append("unsupported_identity_format")
                except (ValueError, OverflowError):
                    issues.append("unsafe_account_number")
        values.append(_text(value))
    nonempty = [value for value in values if value]
    compare = [normalize(value) if normalize else value for value in nonempty]
    if len(set(compare)) > 1:
        issues.append("conflicting_field")
    return nonempty[0] if nonempty else ""


_CARD_HEADERS = frozenset(_norm(value) for value in (
    "本方卡号", "交易卡号", "卡号", "交易对方卡号", "对方卡号", "对手卡号",
    "付款人卡号", "付款卡号", "转出卡号", "收款人卡号", "收款卡号", "转入卡号"))


def _identity_pair(row, candidates, field, issues):
    """Validate each identifier kind; new contracts retain typed overrides."""
    accounts, cards = [], []
    items = list(candidates.get(field, ()))
    if field in {"self_account", "counter_account"}:
        card_items = candidates.get(field.replace("account", "card"), ())
    else:
        card_items = ()
    for item in items:
        (cards if item.get("identity_kind") == "card" or
         _norm(item.get("source_name", "")) in _CARD_HEADERS else accounts).append(item)
    for item in card_items:
        (accounts if item.get("identity_kind") == "account" else cards).append(item)
    return (_field(row, {field: accounts}, field, issues),
            _field(row, {field: cards}, field, issues))


def _amount_field(row, candidates, name, issues):
    found = []
    for candidate in candidates.get(name, []):
        amount, error = _money(_value(row, candidate), candidate)
        if error:
            issues.append("invalid_balance" if name in {"balance", "balance_before"} else error)
        if amount is not None:
            found.append(amount)
    if len(set(found)) > 1:
        issues.append("invalid_balance" if name in {"balance", "balance_before"} else "conflicting_amount")
    return found[0] if found else None


def _reference_field(row, candidates, issues):
    """Separate the source transaction ID from its independent voucher."""
    references = candidates.get("trade_no", [])
    primary = [item for item in references if item.get("reference_role") == "source_transaction_id"]
    vouchers = [item for item in references if item.get("reference_role") == "independent_voucher_reference"]
    transaction = _field(row, {"trade_no": primary}, "trade_no", issues)
    voucher = _field(row, {"trade_no": vouchers}, "trade_no", issues)
    # Keep the already accepted v5 value when the actual transaction ID is blank.
    return (transaction, "source_transaction_id") if transaction else (
        (voucher, "legacy_voucher_fallback") if voucher else ("", "none"))


def _financial_candidates(row, candidates):
    """Apply only explicitly scoped row units; never carry one row into another."""
    resolved, issues = dict(candidates), []
    for field in (*table_financial.MONEY_FIELDS, *table_contracts.MONEY_FIELDS):
        unit_columns = candidates.get(field + "_unit", [])
        if not unit_columns:
            continue
        tokens = [_unit_token(_value(row, item)) for item in unit_columns
                  if _text(_value(row, item))]
        factors = {_UNIT.get(token) for token in tokens}
        items = candidates.get(field, [])
        if field == "amount" and tokens and not items and any(candidates.get(key) for key in ("income_amt", "expense_amt")):
            # A generic amount unit does not declare which split column it
            # describes. Income, expense and balance have their own scopes.
            issues.append("unknown_unit")
        resolved[field] = []
        for item in items:
            explicit = item["unit"] != "元/默认基础单位"
            current = dict(item)
            if _text(_value(row, item)):
                if None in factors:
                    issues.append("unknown_unit")
                elif len(factors) > 1:
                    issues.append("conflicting_unit")
                elif not factors:
                    if not explicit:
                        issues.append("unknown_unit")
                else:
                    factor = next(iter(factors))
                    if explicit and item["factor"] != factor:
                        issues.append("conflicting_unit")
                    elif not explicit:
                        current.update(factor=factor, unit=tokens[0])
                    if any(token in {"人民币元", "cny"} for token in tokens):
                        current["unit_currency"] = "CNY"
            resolved[field].append(current)
    return resolved, issues


def _financial_times(row, candidates, issues):
    """Normalize every date/clock candidate before comparing semantic values."""
    clocks, dates = [], []
    present = False
    for item in candidates.get("time_part", []):
        source = _value(row, item)
        raw = _text(source)
        if raw:
            present = True
            if item.get("time_format") and not table_time_formats.validate(
                    _raw(source), item["time_format"], date_only=_date_only(source)):
                issues.append("invalid_time")
                continue
            value = table_financial.clock(raw)
            if value is None:
                issues.append("invalid_time")
            else:
                clocks.append(value)
    for item in candidates.get("trade_time", []):
        value = _value(row, item)
        raw = _text(value)
        if not raw:
            continue
        present = True
        if item.get("time_format") and not table_time_formats.validate(
                _raw(value), item["time_format"], date_only=_date_only(value)):
            issues.append("invalid_time")
            continue
        clock = table_financial.clock(raw)
        if clock is not None:
            clocks.append(clock)
        else:
            dates.append((value, item.get("time_format") in {"YYYY-MM-DD", "YYYY/MM/DD", "YYYYMMDD"}))
    if len(set(clocks)) > 1:
        issues.append("conflicting_time")
    times, kinds = [], []
    for value, declared_date in dates:
        parsed = _time(value)
        if parsed is None:
            issues.append("invalid_time")
            continue
        only_date = _date_only(value) or declared_date
        if clocks:
            if only_date:
                parsed = dt.datetime.combine(parsed.date(), clocks[0])
            elif parsed.time() != clocks[0]:
                issues.append("conflicting_time")
        times.append(parsed)
        kinds.append(only_date and not clocks)
    if not times:
        issues.append("invalid_time" if present else "missing_time")
    elif len(set(times)) > 1:
        detailed = {value for value, date_only in zip(times, kinds) if not date_only}
        if len(detailed) == 1 and all(value.date() == next(iter(detailed)).date() for value in times):
            times = list(detailed)
        else:
            issues.append("conflicting_time")
    return times


def _interpret(row, candidates, mode, *, reference_policy=None, financial=False, source_contract=None, source_roles=False, identity_values=None):
    issues, notes = [], {}
    if financial:
        candidates, issues = _financial_candidates(row, candidates)
    critical_fields = set(candidates) - {"summary", "channel", "feedback_reason"}
    if any(isinstance(value := _value(row, item), _WorkbookValue) and value.formula for field in critical_fields for item in candidates[field]):
        return None, "unresolved", "formula_value", {"financial_effect": "undetermined", "validation_issues": "formula_value"}
    if (candidates.get("feedback_reason")
            and any(_text(_value(row, item)) for item in candidates["feedback_reason"])
            and not any(_text(_value(row, item)) for field in ("trade_time", "amount", "income_amt", "expense_amt", "direction") for item in candidates.get(field, []))):
        return None, "non_transaction", "feedback_only", {"financial_effect": "not_a_flow"}
    status_candidates = candidates.get("status", [])
    statuses = [_text(_value(row, item)) for item in status_candidates]
    if statuses:
        status_values = {_enum(value) for value in statuses}
        code_basis = all(item.get("source_name") == "交易是否成功" for item in status_candidates)
        positive_statuses = {"成功", "交易成功", "success", "succeeded"} | ({"01", "1"} if code_basis else set())
        if not status_values <= positive_statuses:
            if status_values and status_values <= {"失败", "交易失败", "failed", "failure", "不成功"}:
                return None, "non_transaction", "failed_transaction", {"financial_effect": "not_a_flow", "reported_status": "explicit_failure"}
            issues.append("unknown_status")
        else:
            notes["reported_status"] = "source_reported_success"
    if financial:
        times = _financial_times(row, candidates, issues)
    else:
        times, time_kinds = [], []
        time_parts = [_text(_value(row, candidate)) for candidate in candidates.get("time_part", []) if _text(_value(row, candidate))]
        # A dedicated date and ordinary 'time' column is a supported split format.
        time_candidates = candidates.get("trade_time", [])
        date_values, clock_values = [], []
        for candidate in time_candidates:
            value = _value(row, candidate)
            raw = _text(value)
            if re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?", raw):
                clock_values.append(raw)
            elif raw:
                date_values.append(value)
        time_parts += clock_values
        if len(set(time_parts)) > 1:
            issues.append("conflicting_time")
        for value in date_values:
            parsed = _time(value)
            if parsed is None:
                issues.append("invalid_time")
            elif time_parts:
                clock = _time("2000-01-01 " + time_parts[0])
                if clock is None:
                    issues.append("invalid_time")
                elif _date_only(value):
                    parsed = dt.datetime.combine(parsed.date(), clock.time())
                elif parsed.time() != clock.time():
                    issues.append("conflicting_time")
            if parsed is not None:
                times.append(parsed)
                time_kinds.append(_date_only(value) and not time_parts)
        if not times:
            issues.append("invalid_time" if date_values or time_parts else "missing_time")
        elif len(set(times)) > 1:
            # A date-only candidate may corroborate a full timestamp's calendar day.
            detailed = {value for value, date_only in zip(times, time_kinds) if not date_only}
            if len(detailed) == 1 and all(value.date() == next(iter(detailed)).date() for value in times):
                times = list(detailed)
            else:
                issues.append("conflicting_time")
    amount = _amount_field(row, candidates, "amount", issues)
    income = _amount_field(row, candidates, "income_amt", issues)
    expense = _amount_field(row, candidates, "expense_amt", issues)
    balance = _amount_field(row, candidates, "balance", issues)
    before = _amount_field(row, candidates, "balance_before", issues) if source_contract else None
    if (income is not None and income < 0) or (expense is not None and expense < 0):
        issues.append("negative_amount")
    split_direction, split_amount = None, None
    if income not in (None, 0) and expense not in (None, 0):
        issues.append("conflicting_income_expense")
    elif income not in (None, 0):
        split_direction, split_amount = "in", income
    elif expense not in (None, 0):
        split_direction, split_amount = "out", expense
    elif income is not None or expense is not None:
        split_amount = Decimal("0.00")
    direction_values = []
    for candidate in candidates.get("direction", []):
        value = _enum(_value(row, candidate))
        if value in _IN:
            direction_values.append("in")
        elif value in _OUT:
            direction_values.append("out")
        elif value:
            issues.append("invalid_direction")
    if len(set(direction_values)) > 1:
        issues.append("conflicting_direction")
    direction = direction_values[0] if direction_values else None
    explicit_direction = direction
    signed = any(candidate["signed"] and _money(_value(row, candidate), candidate)[0] is not None for candidate in candidates.get("amount", []))
    sign_direction = "in" if amount is not None and amount > 0 else "out" if amount is not None and amount < 0 else None
    if amount is not None and amount < 0 and (not signed or mode != "B"):
        issues.append("negative_amount")
    if signed and sign_direction:
        if direction and direction != sign_direction:
            issues.append("conflicting_direction")
        direction = direction or sign_direction
        notes["sign_convention"] = "表头明示带符号金额：正收入、负支出"
    if split_direction:
        if direction and direction != split_direction:
            issues.append("conflicting_direction")
        direction = direction or split_direction
    if amount is not None and split_amount is not None and abs(amount) != split_amount:
        issues.append("conflicting_amount")
    amount = amount if amount is not None else split_amount
    if amount is None:
        issues.append("missing_amount")
    if mode == "B" and not direction and amount != 0 and not source_contract:
        issues.append("missing_direction")
    if mode is None:
        issues.append("unknown_table")
    fields, identities = {}, {}
    for field in ("trade_no", "payer_name", "payer_account", "payer_bank", "payee_name", "payee_account", "payee_bank", "self_name", "self_account", "counter_name", "counter_account", "counter_bank"):
        if field == "trade_no" and reference_policy == TRANSACTION_REFERENCE_POLICY:
            fields[field], notes["trade_no_basis"] = _reference_field(row, candidates, issues)
        elif source_roles and field in {"payer_account", "payee_account", "self_account", "counter_account"}:
            identities[field] = _identity_pair(row, candidates, field, issues)
            fields[field] = identities[field][0] or identities[field][1]
        else:
            fields[field] = _field(row, candidates, field, issues)
    for field in ("summary", "channel"):
        fragments = list(dict.fromkeys(_text(_value(row, item)) for item in candidates.get(field, []) if _text(_value(row, item))))
        fields[field] = "；".join(fragments)
    if not fields["self_account"] and candidates.get("self_card") and candidates.get("self_account"):
        # This fallback is bounded to the explicit source pair, not arbitrary cards.
        account_columns = candidates["self_account"]
        if len(account_columns) == 1 and account_columns[0].get("source_name") == "交易账号":
            fields["self_account"] = _field(row, candidates, "self_card", issues)
            notes["source_account_column"] = "交易卡号"
    card_fallbacks = set()
    if source_roles:
        # Validate cards even when an explicit account is present. Only a blank
        # account may use its independently named card; neither overwrites the
        # other and same-role duplicate candidates must still agree.
        for account, card in (("self_account", "self_card"), ("counter_account", "counter_card")):
            account_value, card_value = identities[account]
            if not account_value and card_value:
                fields[account] = card_value
                card_fallbacks.add(account)
                if account == "self_account":
                    notes["source_account_column"] = next(item["source_name"] for item in (
                        *candidates.get(card, ()), *(item for item in candidates.get(account, ())
                                                   if item.get("identity_kind") == "card" or
                                                   _norm(item["source_name"]) in _CARD_HEADERS))
                        if _text(_value(row, item)))
    if fields["self_account"] and "source_account_column" not in notes:
        owner_accounts = [item for item in candidates.get("self_account", ())
                          if item.get("contract_owner") and item.get("identity_kind") == "account"
                          and _text(_value(row, item))]
        if owner_accounts:
            notes["source_account_column"] = owner_accounts[0]["source_name"]
    if fields["self_account"] and "source_account_column" not in notes:
        notes["source_account_column"] = "mapped_self_account"
    for field in ("payer_account", "payee_account", "self_account", "self_card", "counter_account", "counter_card"):
        for candidate in candidates.get(field, []):
            value = _raw(_value(row, candidate))
            if isinstance(value, (int, float)) and not isinstance(value, bool) and len(str(abs(int(value)))) > 15:
                issues.append("unsafe_account_number")
    if mode == "A" and not ((fields["payer_account"] or fields["payer_name"]) and (fields["payee_account"] or fields["payee_name"])):
        issues.append("missing_identity")
    elif not fields["self_account"] and (mode == "B" or any(
            item.get("contract_owner") for field in ("self_account", "self_card")
            for item in candidates.get(field, ()))):
        # A uniform source's owner is a required source fact in either layout.
        # Switching to party mode cannot turn a reference into the missing owner.
        issues.append("missing_identity")
    if source_contract:
        # All semantic candidates remain active; the contract was compiled once
        # for the entire region, not selected from each row's balance fit.
        after = _amount_field(row, {"balance": [item for item in candidates.get("balance", [])
                                              if item.get("same_entry_after")]}, "balance", issues)
        paired_before = before if any(item.get("same_entry_before") and _text(_value(row, item))
                                      for item in candidates.get("balance_before", [])) else None
        party_direction = None
        if mode == "A" and fields["self_account"] and "self_account" not in card_fallbacks:
            own = fields["self_account"]
            payer_account = identities["payer_account"][0] if source_roles else fields["payer_account"]
            payee_account = identities["payee_account"][0] if source_roles else fields["payee_account"]
            if own == payer_account and own != payee_account:
                party_direction = "out"
            elif own == payee_account and own != payer_account:
                party_direction = "in"
        direction, contract_issues, contract_notes = table_contracts.resolve_direction(
            source_contract, mode="B" if fields["self_account"] else mode, amount=amount, direction=explicit_direction,
            signed_direction=sign_direction if signed else None, split_direction=split_direction, party_direction=party_direction,
            before=paired_before, after=after)
        issues.extend(contract_issues)
        notes.update(contract_notes)
        if mode == "B" and not direction and amount != 0:
            issues.append("missing_direction")
    if source_roles and direction:
        # A card can identify its account ledger without being equal to the
        # bank account number. A known direction attaches an explicitly named
        # party account to the same role; it never derives direction from the
        # equality or difference of two identifier kinds.
        own_side, counter_side = ("payee", "payer") if direction == "in" else ("payer", "payee")
        for account, side in (("self_account", own_side), ("counter_account", counter_side)):
            field = side + "_account"
            own_pair, party_pair = identities[account], identities[field]
            if not any(own_pair):
                continue
            account_conflict = bool(own_pair[0] and party_pair[0] and own_pair[0] != party_pair[0])
            own_direction_conflict = account_conflict and account == "self_account" and mode == "A"
            if own_direction_conflict:
                issues.append("conflicting_direction")
            if (account_conflict and not own_direction_conflict
                    or own_pair[1] and party_pair[1] and own_pair[1] != party_pair[1]):
                # A new same-kind conflict must not be hidden by a missing
                # identity, while earlier financial errors retain priority.
                position = next((index for index, issue in enumerate(issues)
                                 if issue in {"missing_identity", "conflicting_direction"}), len(issues))
                issues.insert(position, "conflicting_field")
            combined = tuple(own or party for own, party in zip(own_pair, party_pair))
            identities[account] = identities[field] = combined
            fields[account] = combined[0] or combined[1]
            if any(party_pair) and not own_direction_conflict:
                fields[field] = fields[account]
            if combined[0]:
                card_fallbacks.discard(account)
                if account == "self_account" and not own_pair[0]:
                    notes["source_account_column"] = next(item["source_name"] for item in candidates[field]
                        if _norm(item["source_name"]) not in _CARD_HEADERS and _text(_value(row, item)))
    if mode == "A" and fields["self_account"] and direction and "self_account" not in card_fallbacks:
        expected = fields["payee_account"] if direction == "in" else fields["payer_account"]
        if fields["self_account"] != expected:
            issues.append("conflicting_direction")
    currencies = []
    for candidate in candidates.get("currency", []):
        value = _text(_value(row, candidate))
        if not value:
            continue
        currency = _CURRENCY.get(value, value.upper())
        if currency not in _ISO_CURRENCIES:
            issues.append("invalid_currency")
        else:
            currencies.append(currency)
    if len(set(currencies)) > 1:
        issues.append("invalid_currency")
    currency = currencies[0] if currencies else "CNY"
    unit_currencies = {candidate.get("unit_currency") for name in (*table_financial.MONEY_FIELDS, *(table_contracts.MONEY_FIELDS if source_contract else ())) for candidate in candidates.get(name, []) if candidate.get("unit_currency") and _text(_value(row, candidate))}
    if unit_currencies and unit_currencies != {currency}:
        issues.append("currency_unit_conflict")
    if not currencies:
        notes["currency_basis"] = "沿用普通表格历史默认币种CNY；来源未提供币种"
    if mode == "B":
        own = (fields["self_name"], fields["self_account"], "")
        counter = (fields["counter_name"], fields["counter_account"], fields["counter_bank"])
        payer, payee = (counter, own) if direction == "in" else (own, counter)
        fields.update(payer_name=payer[0], payer_account=payer[1], payer_bank=payer[2], payee_name=payee[0], payee_account=payee[1], payee_bank=payee[2])
        if source_roles:
            own_side, counter_side = ("payee", "payer") if direction == "in" else ("payer", "payee")
            identities[own_side + "_account"] = identities["self_account"]
            identities[counter_side + "_account"] = identities["counter_account"]
        if not fields["counter_account"]:
            notes["counterparty_basis"] = "来源未提供对手账号；保留为空，不以姓名代造账号"
    if source_roles and identity_values is not None:
        identity_values.update(identities)
    issues = list(dict.fromkeys(issues))
    if issues:
        return None, "unresolved", issues[0], {**notes, "validation_issues": ",".join(issues), "financial_effect": "undetermined"}
    if amount == 0:
        return None, "non_monetary", "zero_amount", {**notes, "financial_effect": "no_movement"}
    result = {key: fields[key] for key in ("trade_no", "payer_name", "payer_account", "payer_bank", "payee_name", "payee_account", "payee_bank", "self_name", "self_account", "summary", "channel")}
    result.update(trade_time=times[0], amount=abs(amount), balance=balance, currency=currency)
    return result, "accepted", "transaction_valid", {**notes, "financial_effect": "source_reported", "flow_sha256": flow_digest(result)}


def _role_conflict(outcome):
    notes = dict(outcome[3])
    issues = list(filter(None, notes.get("validation_issues", "").split(",")))
    primary = outcome[2] if issues and set(issues) - {"missing_identity", "conflicting_direction"} else "conflicting_field"
    notes.pop("flow_sha256", None)
    notes.update(financial_effect="undetermined", counterparty_basis=ROLE_CONFLICT_BASIS,
                 validation_issues=",".join(dict.fromkeys([primary, *issues, "conflicting_field"])))
    return None, "unresolved", primary, notes


def _interpret_candidates(row, candidates, mode, *, reference_policy=None, financial=False, source_contract=None, source_roles=False):
    """Resolve explicit role layouts, never rank them by row acceptance rate.

    Financial validation stays in the established interpreter. Only missing
    identities can be completed from another explicitly directed layout.
    """
    selected_identities = {}
    selected = _interpret(row, candidates, mode, reference_policy=reference_policy, financial=financial, source_contract=source_contract, source_roles=source_roles, identity_values=selected_identities)
    party_fields = tuple(side + "_" + field for side in ("payer", "payee") for field in ("name", "account", "bank"))
    if selected[1] == "unresolved":
        identity_issues = []
        for field in (*party_fields, "self_name", "self_account", "self_card", "counter_name", "counter_account", "counter_card", "counter_bank"):
            if len(candidates.get(field, [])) > 1:
                if source_roles and field in {"payer_account", "payee_account", "self_account", "counter_account"}:
                    _identity_pair(row, candidates, field, identity_issues)
                else:
                    _field(row, candidates, field, identity_issues)
        if "conflicting_field" in identity_issues:
            return _role_conflict(selected)
    has_owner = candidates.get("self_account") or source_roles and candidates.get("self_card")
    has_counter = any(candidates.get("counter_" + field) for field in ("name", "account", "bank")) or source_roles and candidates.get("counter_card")
    if (not has_owner or not has_counter and not (source_roles and candidates.get("self_card"))
            or not any(candidates.get(field) for field in party_fields)
            or selected[1] in {"non_transaction", "non_monetary"}):
        return selected
    peer_identities = selected_identities if mode == "B" else {}
    peer = selected if mode == "B" else _interpret(row, candidates, "B", reference_policy=reference_policy, financial=financial, source_contract=source_contract, source_roles=source_roles, identity_values=peer_identities)
    if peer[0] is not None:
        merged = dict(peer[0])
        for field in party_fields:
            if source_roles and field in peer_identities:
                explicit = _identity_pair(row, candidates, field, [])
                projected = peer_identities[field]
                if any(left and right and left != right for left, right in zip(explicit, projected)):
                    return _role_conflict(selected)
                merged[field] = projected[0] or explicit[0] or projected[1] or explicit[1]
                continue
            value = _field(row, candidates, field, [])
            if value and merged[field] and value != merged[field]:
                return _role_conflict(selected)
            if value:
                merged[field] = value
        # A missing source identity is the only repair permitted here. In
        # particular, signed negatives rejected in A do not become B successes.
        issues = set(filter(None, selected[3].get("validation_issues", "").split(",")))
        if selected[0] is not None or selected[1] == "unresolved" and issues and issues <= {"missing_identity", "conflicting_direction"}:
            return merged, "accepted", "transaction_valid", {**peer[3], "flow_sha256": flow_digest(merged)}
        return selected
    if selected[0] is None or mode != "A":
        return selected
    # A already names both parties. Test every possible orientation of sparse
    # B evidence; if more than one remains, keep the explicit A facts. Do not
    # infer a missing direction for a B-only ledger.
    own = {"name": selected[0]["self_name"], **({"account": selected[0]["self_account"]} if not source_roles else {})}
    counter = {field: _field(row, candidates, "counter_" + field, []) for field in (
        ("name", "bank") if source_roles else ("name", "account", "bank"))}
    directions = set()
    for item in candidates.get("direction", []):
        value = _enum(_value(row, item))
        if value in _IN | _OUT:
            directions.add("in" if value in _IN else "out")
    for field, direction in (("income_amt", "in"), ("expense_amt", "out")):
        amount = _amount_field(row, candidates, field, [])
        if amount is not None and amount > 0:
            directions.add(direction)
    for item in candidates.get("amount", []):
        if item["signed"]:
            amount, error = _money(_value(row, item), item)
            if not error and amount:
                directions.add("in" if amount > 0 else "out")
    for direction in directions or {"in", "out"}:
        own_side, counter_side = ("payee", "payer") if direction == "in" else ("payer", "payee")
        if source_roles:
            pairs = ((selected_identities["self_account"], selected_identities[own_side + "_account"]),
                     (selected_identities["counter_account"], selected_identities[counter_side + "_account"]))
            if (all(not left or not right or left == right for expected, actual in pairs
                    for left, right in zip(expected, actual))
                    and all(not value or not selected[0][side + "_" + field] or value == selected[0][side + "_" + field]
                            for side, values in ((own_side, own), (counter_side, counter))
                            for field, value in values.items() if field != "account")):
                return selected
            continue
        if all(not value or not selected[0][side + "_" + field] or value == selected[0][side + "_" + field]
               for side, fields in ((own_side, own), (counter_side, counter)) for field, value in fields.items()):
            return selected
    return _role_conflict(selected)


def _region(headers, mode, mapping, *, strict, reference_policy=None, profile=IDENTITY_PROFILE):
    """Resolve labels against this region's positions, never an earlier region."""
    candidates = _candidates(headers, profile)
    source_role = table_source_roles.classify(
        [_text(header) for header in headers], query_conditions=profile in QUERY_OWNER_PROFILES
    ) if profile in SOURCE_ROLE_PROFILES else None
    query_owner = profile in QUERY_OWNER_PROFILES and source_role in {
        "transaction_explicit", "transaction_query_only"}
    labels = _columns(headers)
    for items in candidates.values():
        for candidate in items:
            candidate["source_name"] = _text(headers[candidate["index"]])
    error = None
    pinned_mapping = {}
    for field, column in (mapping or {}).items():
        if field not in (mapping_fields(profile) if profile in CONTRACT_PROFILES else FINANCIAL_MAPPING_FIELDS if profile == FINANCIAL_PROFILE else _ALIASES) or not isinstance(column, str):
            raise TableImportError("表格映射字段无效")
        matches = [index for index, label in enumerate(labels) if label == column]
        if len(matches) != 1:
            if strict:
                raise TableImportError("表格映射列不存在或不唯一")
            error = "region_mapping_conflict"
            continue
        index = matches[0]
        if query_owner and table_source_roles.query_condition(_text(headers[index])):
            if strict:
                raise TableImportError("查询条件注记不是交易事实字段")
            error = "region_mapping_conflict"
            continue
        identity = next((item for items in candidates.values() for item in items
                         if item["index"] == index and item.get("contract_owner")), None)
        if query_owner and (
                field in {"self_account", "self_card"} and identity is None
                or identity is not None and (field not in {"self_account", "self_card", "payer_account", "payee_account"}
                                             or field == "self_card" and identity["identity_kind"] != "card")):
            if strict:
                raise TableImportError("统一交易明细的本方映射必须保留本方账号、卡号职责")
            error = "region_mapping_conflict"
            continue
        pinned_mapping[field] = column
        base, factor, unit = _header_parts(headers[index], profile)
        explicit = {"index": index, "factor": factor, "unit": unit, "signed": base in _SIGNED, "source_name": _text(headers[index]), "unit_currency": "CNY" if unit in {"人民币元", "cny"} else ""}
        if query_owner and identity:
            explicit.update(identity_kind=identity["identity_kind"], contract_owner=True)
        declaration = _declared_time_header(headers[index], profile)
        if declaration and field in {"trade_time", "time_part"}:
            explicit["time_format"] = declaration[1]
        if profile in FINANCIAL_PROFILES and field in {
                *table_financial.MONEY_FIELDS, "trade_time", "time_part", "direction",
                "currency", "status", "amount_unit", "income_amt_unit", "expense_amt_unit", "balance_unit"}:
            explicit["basis"] = "explicit_financial_header"
        if profile in CONTRACT_PROFILES:
            explicit.update(same_entry_after=field == "balance" and base in _AFTER_HEADERS,
                            same_entry_before=field == "balance_before" and base in _BEFORE_HEADERS)
            if field in {"balance_before", "balance_before_unit"}:
                explicit["basis"] = "explicit_source_contract"
        # An override still cannot suppress other monetary alias candidates.
        if (query_owner and field in {"self_account", "self_card"}) or field in ({"amount", "income_amt", "expense_amt", "balance"} |
                     ({"balance_before", "balance_before_unit"} if profile in CONTRACT_PROFILES else set()) |
                     ({"trade_time", "time_part", "direction", "currency", "status", "amount_unit", "income_amt_unit", "expense_amt_unit", "balance_unit"} if profile in FINANCIAL_PROFILES else set())):
            if index not in {item["index"] for item in candidates.get(field, [])}:
                candidates.setdefault(field, []).append(explicit)
        else:
            candidates[field] = [explicit]
    actual_mode = mode or _shape(candidates, profile)
    if not actual_mode and source_role == "transaction_query_only":
        actual_mode = "B"
    current_mapping = {field: labels[items[0]["index"]] for field, items in candidates.items() if len(items) == 1}
    if profile in FINANCIAL_PROFILES:
        # Multiple candidates remain active validators, while the explicit
        # choice must survive repreview/commit so custom columns are not lost.
        current_mapping.update(pinned_mapping)
    active_reference_policy = None
    if ((reference_policy == TRANSACTION_REFERENCE_POLICY
            and len(headers) == len(BANK_FEEDBACK_32_COLUMNS)
            and frozenset(_text(header) for header in headers) == BANK_FEEDBACK_32_COLUMNS)
            or source_role in {"transaction_explicit", "transaction_query_only"} and not (mapping and "trade_no" in mapping)):
        active_reference_policy = TRANSACTION_REFERENCE_POLICY
        for item in candidates.get("trade_no", []):
            item["reference_role"] = ("source_transaction_id" if item["source_name"] == "交易流水号"
                                      else "independent_voucher_reference")
        current_mapping["trade_no"] = labels[next(item["index"] for item in candidates["trade_no"]
                                                   if item["source_name"] == "交易流水号")]
    return candidates, labels, actual_mode, current_mapping, error, active_reference_policy


def _later_header(row, candidates, profile=IDENTITY_PROFILE):
    # Keep the fast path for ordinary dated financial rows. A complete new
    # header may put a date-named extra column at the previous time coordinate;
    # v16 checks that structure only when the old monetary fields cannot parse.
    has_old_time = any(_time(_value(row, item)) is not None for item in candidates.get("trade_time", []))
    if has_old_time:
        if profile != REORDERED_TIME_PROFILE or any(
                _money(_value(row, item), item)[0] is not None
                for field in ("amount", "income_amt", "expense_amt")
                for item in candidates.get(field, [])):
            return False
    monetary_fields = ("amount", "income_amt", "expense_amt", "balance")
    has_old_money = any(_money(_value(row, item), item)[0] is not None
                        for field in monetary_fields for item in candidates.get(field, []))
    if has_old_money and profile not in REGION_HEADER_PROFILES:
        return False
    time_aliases = _ALIASES["trade_time"] | (_FINANCIAL_ALIASES["trade_time"] if profile in FINANCIAL_PROFILES else set())
    if not any(_norm(value) in time_aliases or _declared_time_header(value, profile) for item in row
               if (value := _text(item)) and not value[0].isdigit() and len(value) <= 128):
        return False
    found = _candidates(row, profile)
    money_fields = ("amount", "income_amt", "expense_amt")
    if not found.get("trade_time") or not any(found.get(key) for key in money_fields):
        return False
    if _shape(found, profile):
        return True
    # A complete new schema can move the amount column and put a numeric
    # unrelated header at its old coordinate. It must take precedence over
    # that old coordinate's monetary value. Partial label-like data still
    # cannot revoke a region when its current monetary cell is populated.
    if has_old_time or has_old_money:
        return False
    # A partial schema can revoke the old interpretation only when its actual
    # time and money coordinates contain those header labels. Words such as
    # "date" in an account and "amount" in a memo are not a regional boundary,
    # including when both the real time and amount cells are empty.
    old_time = {item["index"] for item in candidates.get("trade_time", [])}
    old_money = {item["index"] for key in money_fields for item in candidates.get(key, [])}
    new_time = {item["index"] for item in found["trade_time"]}
    new_money = {item["index"] for key in money_fields for item in found.get(key, [])}
    return bool(old_time & new_time and old_money & new_money)


def _region_evidence(headers, candidates, labels, path, header_row):
    return [{"field": "voucher_reference" if item.get("reference_role") == "independent_voucher_reference" else field,
             "column": labels[item["index"]], "column_index": item["index"] + 1,
             "unit": item["unit"], "unit_factor": str(item["factor"]) if item["factor"] is not None else "",
             "candidate_count": sum(other.get("reference_role") == item["reference_role"] for other in items) if "reference_role" in item else len(items),
             "basis": item.get("reference_role", item.get("basis", "exact_header_alias")), "member_path": path,
             "header_row": header_row,
             **({"time_format": item["time_format"]} if "time_format" in item else {})}
            for field, items in candidates.items() for item in items]


def _append_auxiliary_sheet(plan, sheet):
    rows = sheet["rows"]
    header_no = next((index for index, row in enumerate(rows) if any(_text(value) for value in row)), None)
    if header_no is None:
        return False
    headers = rows[header_no]
    role = table_source_roles.classify([_text(value) for value in headers])
    if role not in {"account", "client", "subaccount"}:
        return False
    # A later transaction region must go through the ordinary complete parser.
    if _locate_header(rows[header_no + 1:], plan["profile"])[0] is not None:
        return False
    width = len(headers)
    meaningful_width = max(index + 1 for index, value in enumerate(headers) if _text(value))
    contract = table_contracts.compile_contract({})
    labels = _columns(headers)
    plan["documents"].append({"member_path": sheet["path"], "sheet": sheet["name"],
                              "header_row": header_no + 1, "header_rows": [header_no + 1],
                              "columns": labels, "mapping": {}, "mode": "", "source_role": role,
                              "source_rows": len(rows), "row_count": len(rows) - 1})
    for index, row in enumerate(rows, start=1):
        if index == header_no + 1:
            continue
        valid = len(row) == width and not any(_text(value) for value in row[meaningful_width:])
        reason = "auxiliary_table" if valid else "column_count_mismatch"
        details = {"sheet": sheet["name"], "source_row": str(index), "source_role": role,
                   "financial_effect": "not_a_flow" if valid else "undetermined",
                   "reason": REASONS[reason], **contract.details()}
        plan["records"].append({"record_no": len(plan["records"]) + 1, "member_path": sheet["path"],
                                "source_row": index, "record_key": record_key(sheet["path"], sheet["sha256"], index),
                                "disposition": "auxiliary" if valid else "unresolved", "reason_code": reason,
                                "details": details})
    plan["members"].append({"member_path": sheet["path"], "role": "auxiliary", "sha256": sheet["sha256"],
                            "size_bytes": sheet["size"], "encoding": sheet["encoding"], "row_count": len(rows) - 1})
    return True


def parse(source, *, suffix, mode=None, mapping=None, reference_policy=None, profile=PROFILE):
    # New exact cross-field arithmetic is independent of ambient Decimal
    # precision. Historical parsers keep their original behavior and receipts.
    if profile in CONTRACT_PROFILES:
        with localcontext() as context:
            context.prec = max(context.prec, 50)
            return _parse(source, suffix=suffix, mode=mode, mapping=mapping,
                          reference_policy=reference_policy, profile=profile)
    return _parse(source, suffix=suffix, mode=mode, mapping=mapping,
                  reference_policy=reference_policy, profile=profile)


def _parse(source, *, suffix, mode=None, mapping=None, reference_policy=None, profile=PROFILE):
    """Interpret every source record and return a write-free versioned plan."""
    suffix = suffix.lower()
    if profile not in SUPPORTED_PROFILES:
        raise TableImportError("不支持的表格解析规则")
    if suffix not in {".csv", ".tsv", ".txt", ".xlsx", ".xls"}:
        raise TableImportError("表格文件格式不受支持")
    if mode not in (None, "A", "B") or (mapping is not None and not isinstance(mapping, dict)):
        raise TableImportError("表格映射参数无效")
    if reference_policy is not None and reference_policy != TRANSACTION_REFERENCE_POLICY:
        raise TableImportError("表格引用策略无效")
    if reference_policy is not None and mapping and "trade_no" in mapping:
        raise TableImportError("表格引用策略不能与交易号映射同时使用")
    raw = _bytes(source)
    sheets = _excel(raw, suffix) if suffix in {".xlsx", ".xls"} else _csv(raw, suffix, profile)
    plan = {"profile": profile, "mode": "B", "mapping": {}, "columns": [], "sheet": "workbook" if suffix in {".xlsx", ".xls"} else "", "rows": [], "records": [], "members": [], "documents": [], "mapping_evidence": []}
    selected = False
    for sheet in sheets:
        raw_rows = sheet["rows"]
        if profile in SOURCE_ROLE_PROFILES and _append_auxiliary_sheet(plan, sheet):
            continue
        header_no, _, _ = _locate_header(raw_rows, profile)
        headers = raw_rows[header_no] if header_no is not None else []
        candidates, labels, actual_mode, current_mapping, region_error, active_reference_policy = _region(
            headers, mode, mapping, strict=True, reference_policy=reference_policy, profile=profile)
        source_contract = table_contracts.compile_contract(candidates) if profile in CONTRACT_PROFILES else None
        if not selected and header_no is not None:
            plan.update(mode=actual_mode or "B", mapping=current_mapping, columns=labels)
            selected = True
        plan["mapping_evidence"].extend(_region_evidence(headers, candidates, labels, sheet["path"], header_no + 1 if header_no is not None else None))
        document = {"member_path": sheet["path"], "sheet": sheet["name"], "header_row": header_no + 1 if header_no is not None else None, "header_rows": [header_no + 1] if header_no is not None else [], "columns": labels, "mapping": current_mapping, "mode": actual_mode or "", "source_rows": len(raw_rows), "row_count": len(raw_rows) - (1 if header_no is not None else 0)}
        if source_contract:
            document.update(source_contract=table_contracts.VERSION, contract_sha256=source_contract.digest,
                            amount_model=source_contract.amount_model, candidate_scope="all_region_candidates")
        plan["documents"].append(document)
        member_count = 0
        for index, row in enumerate(raw_rows, start=1):
            if header_no is not None and index == header_no + 1:
                continue
            number = len(plan["records"]) + 1
            key = record_key(sheet["path"], sheet["sha256"], index)
            details = {"sheet": sheet["name"], "source_row": str(index), "financial_effect": "not_a_flow"}
            result = None
            if not any(_text(value) for value in row):
                disposition, reason = "non_transaction", "blank_record"
            elif header_no is not None and index < header_no + 1:
                disposition, reason = "auxiliary", "preamble"
            elif header_no is None:
                disposition, reason = "unresolved", "unknown_table"
                details["financial_effect"] = "undetermined"
            elif _later_header(row, candidates, profile):
                headers = row
                candidates, labels, actual_mode, current_mapping, region_error, active_reference_policy = _region(
                    headers, mode, mapping, strict=False, reference_policy=reference_policy, profile=profile)
                source_contract = table_contracts.compile_contract(candidates) if profile in CONTRACT_PROFILES else None
                document["header_rows"].append(index)
                plan["mapping_evidence"].extend(_region_evidence(headers, candidates, labels, sheet["path"], index))
                disposition, reason = "non_transaction", "repeated_header"
            elif len(row) != len(headers):
                disposition, reason = "unresolved", "row_width"
                details["financial_effect"] = "undetermined"
            elif region_error:
                disposition, reason = "unresolved", region_error
                details.update(financial_effect="undetermined", validation_issues=region_error)
            else:
                interpreter = _interpret_candidates if profile in ROLE_PROFILES else _interpret
                result, disposition, reason, parsed_details = interpreter(
                    row, candidates, actual_mode, reference_policy=active_reference_policy, financial=profile in FINANCIAL_PROFILES, source_contract=source_contract, source_roles=profile in SOURCE_ROLE_PROFILES)
                details.update(parsed_details)
                if source_contract:
                    details.update(source_contract.details())
                if (profile in FINANCIAL_PROFILES and disposition == "unresolved"
                        and FINANCIAL_VALIDATION_ISSUES.intersection(parsed_details.get("validation_issues", "").split(","))
                        and any(item.get("basis") in {"explicit_financial_header", "explicit_source_contract", "explicit_time_format"} for items in candidates.values() for item in items)):
                    details["validation_basis"] = FINANCIAL_CONFLICT_BASIS
            details["reason"] = REASONS[reason]
            if source_contract and reason == "missing_direction":
                details["reason"] = ("缺少明确方向及同一笔交易的完整前后账面余额；普通相邻余额未证明连续性，"
                                     "不能补方向。该笔未计入资金统计和节点关系，收支及关联分析可能不完整")
            if profile in SOURCE_ROLE_PROFILES and reason == "missing_identity":
                details["reason"] = (("统一交易明细的本方账号、卡号及查询帐号、查询卡号均缺失。"
                                      if profile in QUERY_OWNER_PROFILES and any(
                                          item.get("contract_owner") for items in candidates.values()
                                          for item in items)
                                      else "缺少银行返回的本方账号或卡号；查询条件不证明本笔交易主体。")
                                     + "该笔未计入资金统计和节点关系，收支及关联分析可能不完整")
            record = {"record_no": number, "member_path": sheet["path"], "source_row": index, "record_key": key, "disposition": disposition, "reason_code": reason, "details": details}
            plan["records"].append(record)
            if result is not None:
                result.update(_source_row=number, _source_sheet=plan["sheet"], _source_row_kind=ROW_KIND, _record_key=key)
                plan["rows"].append(result)
            member_count += 1
        plan["members"].append({"member_path": sheet["path"], "role": "table", "sha256": sheet["sha256"], "size_bytes": sheet["size"], "encoding": sheet["encoding"], "row_count": member_count})
    return refresh_stats(plan)


def refresh_stats(plan):
    """Recount the shared plan after a transaction-bound history check."""
    counts = {key: 0 for key in ("accepted", "unresolved", "non_monetary", "non_transaction", "auxiliary")}
    reasons, examples, totals = {}, [], {}
    for record in plan["records"]:
        counts[record["disposition"]] += 1
        reasons[record["reason_code"]] = reasons.get(record["reason_code"], 0) + 1
        if record["disposition"] == "unresolved" and len(examples) < 3:
            examples.append({"row": "来源记录", "source_row": record["source_row"], "member_path": record["member_path"], "reason": record["details"].get("reason") or REASONS[record["reason_code"]]})
    for row in plan["rows"]:
        totals[row["currency"]] = totals.get(row["currency"], Decimal(0)) + row["amount"]
    plan["stats"] = {"total": len(plan["records"]), "ok": len(plan["rows"]), "skipped": len(plan["records"]) - len(plan["rows"]), "bad_examples": examples, **counts, "reason_counts": reasons, "currency_totals": {key: format(value, ".2f") for key, value in sorted(totals.items())}}
    return plan
